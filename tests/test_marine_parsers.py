import json
import os
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from marine_kb.parsers import parse_file, _text_quality, _recover_pdf_ocr, _env_bool, _is_word_ooxml


class MarineParsersTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)

    def tearDown(self):
        self.tmp.cleanup()

    def test_empty_and_lock(self):
        path = self.root / 'empty.xlsx'
        path.touch()
        self.assertEqual(parse_file(path)['status'], 'empty')
        locked = self.root / '~$report.docx'
        locked.write_bytes(b'not a document')
        self.assertEqual(parse_file(locked)['status'], 'unsupported')

    def test_plain_text_with_source_locator(self):
        path = self.root / 'report.txt'
        path.write_text('2019年利奇马。\n海门最大增水为3.12米。', encoding='utf-8')
        result = parse_file(path)
        self.assertEqual(result['status'], 'parsed')
        self.assertIn('lines:1-2', result['units'][0]['locator'])
        self.assertIn('3.12', result['units'][0]['text'])

    def test_garbled_pdf_quality(self):
        self.assertEqual(_text_quality('/G0052/G0072/G0031' * 10), 'garbled')
        self.assertEqual(_text_quality('浙江省海洋监测预报中心发布风暴潮警报，具体时效和警报区域以正式发布的产品为准。'), 'usable')

    def test_geojson_is_structured_not_rag(self):
        path = self.root / 'coast.json'
        path.write_text(json.dumps({'type': 'FeatureCollection', 'features': [{'type': 'Feature', 'properties': {'cid': 1, 'RED': '500'}, 'geometry': {'type': 'LineString', 'coordinates': [[120, 28], [121, 29]]}}]}), encoding='utf-8')
        result = parse_file(path)
        self.assertEqual(result['status'], 'parsed')
        self.assertEqual(result['units'], [])
        self.assertEqual(result['metadata']['ingest_mode'], 'structured_only')
        self.assertEqual(result['metadata']['bbox'], [120, 28, 121, 29])
        self.assertFalse(result['metadata']['threshold_unit_confirmed'])

    def test_json_pointer_escaped(self):
        path = self.root / 'knowledge.json'
        path.write_text('{"a/b~c":"原始资料"}', encoding='utf-8')
        self.assertEqual(parse_file(path)['units'][0]['metadata']['json_pointer'], '/a~1b~0c')

    def test_xlsx_only_metadata_incomplete_time(self):
        from openpyxl import Workbook
        path = self.root / 'Surge_detem_TY0216_2026041310.xlsx'
        book = Workbook()
        sheet = book.active
        sheet.title = '确定性预报surge'
        sheet.append(['time', '海门'])
        sheet.append(['05/14:00', -5])
        sheet.append(['06/14:00', 100])
        book.save(path)
        result = parse_file(path)
        self.assertEqual(result['status'], 'parsed')
        self.assertEqual(result['units'], [])
        metadata = result['metadata']['sheets'][0]
        self.assertEqual(metadata['first_time'], '05/14:00')
        self.assertTrue(metadata['incomplete_time'])
        self.assertEqual(metadata['data_row_count'], 2)
        self.assertIsNone(metadata['unit'])

    def test_docx_order_and_fake_doc_extension(self):
        from docx import Document
        path = self.root / 'report.docx'
        document = Document()
        document.add_heading('一、灾情', 1)
        document.add_paragraph('利奇马的历史资料。')
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = '站点'
        table.cell(0, 1).text = '最大增水'
        table.cell(1, 0).text = '海门'
        table.cell(1, 1).text = '3.12 m'
        document.add_paragraph('后续说明。')
        document.save(path)
        disguised = self.root / 'report.doc'
        disguised.write_bytes(path.read_bytes())
        result = parse_file(disguised)
        self.assertEqual(result['status'], 'parsed')
        self.assertEqual([u['kind'] for u in result['units']], ['text', 'text', 'table', 'text'])
        self.assertTrue(result['metadata']['extension_mismatch'])
        self.assertIn('海门 | 3.12 m', result['units'][2]['text'])

    def test_embedded_zip_in_ole_does_not_masquerade_as_docx(self):
        import zipfile
        path = self.root / 'ole-with-theme.doc'
        path.write_bytes(b'\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1' + b'OLE body')
        with zipfile.ZipFile(path, 'a') as archive:
            archive.writestr('[Content_Types].xml', '<Types/>')
            archive.writestr('theme/theme/theme1.xml', '<theme/>')
        self.assertTrue(zipfile.is_zipfile(path))
        self.assertFalse(_is_word_ooxml(path))
        with patch('marine_kb.parsers._parse_legacy_word', return_value={'units': [], 'metadata': {'format': 'ole'}, 'warnings': [], 'status': 'parsed'}) as legacy:
            result = parse_file(path)
        legacy.assert_called_once_with(path)
        self.assertNotIn('extension_mismatch', result['metadata'])

    def test_blank_pdf_needs_ocr(self):
        from pypdf import PdfWriter
        path = self.root / 'scan.pdf'
        writer = PdfWriter()
        writer.add_blank_page(width=100, height=100)
        writer.write(str(path))
        result = parse_file(path)
        self.assertEqual(result['status'], 'needs_ocr')
        self.assertEqual(result['units'], [])
        self.assertEqual(result['metadata']['quarantined_pages'][0]['page'], 1)

    def test_pptx_slide_and_table_provenance(self):
        from pptx import Presentation
        from pptx.util import Inches
        path = self.root / 'forecast.pptx'
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)).text_frame.text = '海浪预报讨论稿'
        table = slide.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(4), Inches(2)).table
        table.cell(0, 0).text = '区域'
        table.cell(0, 1).text = '预报波高'
        table.cell(1, 0).text = '浙北'
        table.cell(1, 1).text = '3—4米'
        presentation.save(path)
        result = parse_file(path)
        self.assertEqual(result['status'], 'parsed')
        self.assertEqual([u['locator'] for u in result['units']], ['slide:1', 'slide:1/table:1'])

    def test_csv_zero_and_missing_are_not_rewritten(self):
        path = self.root / 'damage.csv'
        path.write_text('指标,数值\n伤亡人数,0\n失踪人数,—\n', encoding='utf-8')
        result = parse_file(path)
        self.assertEqual(result['status'], 'parsed')
        self.assertIn('伤亡人数 | 0', result['units'][0]['text'])
        self.assertIn('失踪人数 | —', result['units'][0]['text'])

    def test_local_word_cache_content_sha(self):
        from docx import Document
        import hashlib
        path = self.root / 'original.doc'
        path.write_bytes(b'original Ole content stand-in')
        cache = self.root / 'cache'
        cache.mkdir()
        document = Document()
        document.add_paragraph('本次实际转换缓存正文')
        digest = hashlib.sha256(path.read_bytes()).hexdigest()
        document.save(cache / (digest + '.docx'))
        with patch.dict(os.environ, {'MARINE_KB_WORD_CACHE': str(cache)}):
            result = parse_file(path)
        self.assertEqual(result['status'], 'parsed')
        self.assertEqual(result['metadata']['conversion'], 'local_word_sha256_content_cache')
        self.assertIn('本次实际转换', result['units'][0]['text'])

    def test_encrypted_pdf_unsupported(self):
        from pypdf import PdfWriter
        path = self.root / 'locked.pdf'
        writer = PdfWriter()
        writer.add_blank_page(width=100, height=100)
        writer.encrypt('do-not-guess')
        writer.write(str(path))
        self.assertEqual(parse_file(path)['status'], 'unsupported')

    def test_ocr_default_disabled(self):
        with patch.dict(os.environ, {}, clear=True):
            self.assertFalse(_env_bool('MARINE_LOCAL_OCR_ENABLED'))

    def test_ocr_missing_optional_dependencies_keeps_quarantine(self):
        path = self.root / 'scan.pdf'
        path.write_bytes(b'not read when dependencies unavailable')
        with patch('marine_kb.parsers._ocr_version', return_value=None):
            units, info, warnings = _recover_pdf_ocr(path, [{'page': 1, 'reason': 'garbled'}])
        self.assertEqual(units, [])
        self.assertEqual(info['recovered_pages'], 0)
        self.assertIn('缺少', warnings[0])

    def test_ocr_cached_page_remains_unreviewed(self):
        import hashlib
        path = self.root / 'fake.pdf'
        path.write_bytes(b'only cache identity is needed in test')
        digest = hashlib.sha256(path.read_bytes()).hexdigest()
        cache = self.root / 'ocr-cache'
        cache.mkdir()
        pipeline = 'unit-test-pipeline'
        key = hashlib.sha256(f'{digest}:4:{pipeline}'.encode()).hexdigest()
        (cache / (key + '.json')).write_text(json.dumps({'source_sha256': digest, 'page': 4, 'pipeline_version': pipeline, 'status': 'ok', 'text': '预计未来二十四小时内海洋灾害可能造成影响，具体防御措施应当参照正式发布的预警产品和应急预案。', 'confidence_mean': .98, 'confidence_min': .81, 'line_count': 2, 'low_confidence_lines': 0}), encoding='utf-8')
        with patch.dict(os.environ, {'MARINE_LOCAL_OCR_CACHE': str(cache)}), patch('marine_kb.parsers._ocr_version', return_value=pipeline):
            units, info, warnings = _recover_pdf_ocr(path, [{'page': 4, 'reason': 'garbled'}])
        self.assertEqual(len(units), 1)
        self.assertTrue(units[0]['metadata']['requires_human_review'])
        self.assertFalse(units[0]['metadata']['approved'])
        self.assertEqual(units[0]['metadata']['method'], 'local_rapidocr')
        self.assertTrue(info['requires_human_review'])

    def test_ocr_low_confidence_page_not_promoted(self):
        import hashlib
        path = self.root / 'uncertain.pdf'
        path.write_bytes(b'low confidence cache test')
        digest = hashlib.sha256(path.read_bytes()).hexdigest()
        cache = self.root / 'ocr-low'
        cache.mkdir()
        pipeline = 'unit-test-pipeline'
        key = hashlib.sha256(f'{digest}:1:{pipeline}'.encode()).hexdigest()
        (cache / (key + '.json')).write_text(json.dumps({'source_sha256': digest, 'page': 1, 'pipeline_version': pipeline, 'status': 'ok', 'text': '数值和单位可能识别错误，此行文本虽然够长但置信度不足，所以仍不能被提升为可用正文证据。', 'confidence_mean': .55}), encoding='utf-8')
        with patch.dict(os.environ, {'MARINE_LOCAL_OCR_CACHE': str(cache)}), patch('marine_kb.parsers._ocr_version', return_value=pipeline):
            units, info, warnings = _recover_pdf_ocr(path, [{'page': 1, 'reason': 'garbled'}])
        self.assertEqual(units, [])
        self.assertEqual(info['failed_pages'][0]['page'], 1)

    def test_legacy_disabled_explicit(self):
        path = self.root / 'old.doc'
        path.write_bytes(b'\xd0\xcf\x11\xe0' + b'not real OLE')
        with patch.dict(os.environ, {'MARINE_KB_DISABLE_WORD_COM': '1'}):
            result = parse_file(path)
        self.assertEqual(result['status'], 'unsupported')

    def test_bad_json_is_error_not_evidence(self):
        path = self.root / 'bad.json'
        path.write_text('{invalid', encoding='utf-8')
        result = parse_file(path)
        self.assertEqual(result['status'], 'error')
        self.assertEqual(result['units'], [])


if __name__ == '__main__':
    unittest.main()

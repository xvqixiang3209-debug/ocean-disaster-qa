"""Create a readable Word compilation from collected actual system answers."""
from collections import Counter, OrderedDict
import argparse
from datetime import datetime
import json
from pathlib import Path
import re
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

APP = Path(__file__).resolve().parents[1]
DATA = APP/'tmp/qa_export_20261015/answers.json'
DESTINATION = APP/'exports/海洋灾害160题问答汇编_20261015.docx'


def set_font(run, size=11, bold=None):
    run.font.name = 'Calibri'
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(0,0,0)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'宋体')
    if bold is not None: run.bold=bold


def inline(paragraph, text, size=11):
    for part in re.split(r'(\*\*[^*]+\*\*|`[^`]+`)',text):
        if not part: continue
        bold=part.startswith('**') and part.endswith('**')
        if bold: part=part[2:-2]
        elif part.startswith('`') and part.endswith('`'): part=part[1:-1]
        run=paragraph.add_run(part)
        set_font(run,size,bold)


def field(paragraph, instruction):
    node=OxmlElement('w:fldSimple');node.set(qn('w:instr'),instruction)
    paragraph._p.append(node)


def table(document, records, widths=None):
    columns=max(len(row) for row in records)
    t=document.add_table(rows=0,cols=columns)
    t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    if widths is None:
        widths=([.85]+[(6.5-.85)/(columns-1)]*(columns-1)) if columns>1 else [6.5]
    for column,width in zip(t.columns,widths): column.width=Inches(width)
    borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        b=OxmlElement('w:'+edge);b.set(qn('w:val'),'single');b.set(qn('w:sz'),'4');b.set(qn('w:color'),'D9D9D9');borders.append(b)
    t._tbl.tblPr.append(borders)
    for index,record in enumerate(records):
        row=t.add_row()
        if index==0:
            repeat=OxmlElement('w:tblHeader');row._tr.get_or_add_trPr().append(repeat)
        if max(len(value) for value in record)<600:
            row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        for col,cell in enumerate(row.cells):
            cell.width=Inches(widths[col]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr=cell._tc.get_or_add_tcPr()
            margins=OxmlElement('w:tcMar')
            for edge,value in [('top',90),('bottom',90),('left',95),('right',95)]:
                node=OxmlElement('w:'+edge);node.set(qn('w:w'),str(value));node.set(qn('w:type'),'dxa');margins.append(node)
            tcpr.append(margins)
            fill=OxmlElement('w:shd');fill.set(qn('w:fill'),'203A4D' if index==0 else 'F4F6F8' if index%2==0 else 'FFFFFF');tcpr.append(fill)
            p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.1
            if col>0 and len(record[col] if col<len(record) else '')<=6: p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            inline(p,record[col] if col<len(record) else '',10.5)
            if index==0:
                for run in p.runs: run.bold=True;run.font.color.rgb=RGBColor(255,255,255)
    document.add_paragraph().paragraph_format.space_after=Pt(0)
    return t


def add_answer(document,text):
    lines=text.replace('\r','').split('\n');index=0
    while index<len(lines):
        line=lines[index].strip()
        if not line: index+=1;continue
        if '|' in line and index+1<len(lines) and re.match(r'^\s*\|?\s*:?-{3,}.*\|',lines[index+1]):
            def cells(value): return [v.strip() for v in value.strip().strip('|').split('|')]
            records=[cells(line)];index+=2
            while index<len(lines) and lines[index].strip() and '|' in lines[index]:
                records.append(cells(lines[index]));index+=1
            table(document,records)
            continue
        if re.match(r'^#{1,6}\s',line):
            p=document.add_paragraph();p.paragraph_format.keep_with_next=True
            inline(p,re.sub(r'^#{1,6}\s+','',line))
            for run in p.runs: run.bold=True
        elif re.match(r'^[-*]\s+',line):
            p=document.add_paragraph(style='List Bullet');inline(p,re.sub(r'^[-*]\s+','',line))
        else:
            p=document.add_paragraph();inline(p,line)
        index+=1


def main():
    global DATA, DESTINATION
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--input',type=Path,default=DATA)
    parser.add_argument('--output',type=Path,default=DESTINATION)
    args=parser.parse_args()
    DATA,DESTINATION=args.input,args.output
    report=json.loads(DATA.read_text(encoding='utf-8'))
    assert report['question_count']==160 and len(report['items'])==160
    groups=OrderedDict()
    for item in report['items']: groups.setdefault(item['category'],[]).append(item)
    doc=Document();section=doc.sections[0]
    section.page_width=Inches(8.5);section.page_height=Inches(11)
    section.top_margin=section.bottom_margin=Inches(.8)
    section.left_margin=section.right_margin=Inches(1)
    section.header_distance=section.footer_distance=Inches(.35)
    for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','List Bullet','List Number']:
        style=doc.styles[name];style.font.name='Calibri';style.font.color.rgb=RGBColor(0,0,0)
        style.font.italic=False
        style._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'宋体')
    normal=doc.styles['Normal'];normal.font.size=Pt(11)
    normal.paragraph_format.line_spacing=1.15
    normal.paragraph_format.space_after=Pt(5)
    normal.paragraph_format.widow_control=True
    for name,size,before,after in [('Title',23,0,12),('Subtitle',11,0,12),('Heading 1',16,0,12),('Heading 2',12,12,6)]:
        style=doc.styles[name];style.font.size=Pt(size);style.font.bold=name!='Subtitle'
        style.paragraph_format.space_before=Pt(before);style.paragraph_format.space_after=Pt(after)
        style.paragraph_format.keep_with_next=True
        style._element.get_or_add_pPr().remove(style._element.pPr.find(qn('w:pBdr'))) if style._element.pPr is not None and style._element.pPr.find(qn('w:pBdr')) is not None else None
    doc.add_paragraph('海洋灾害160题与系统回答',style='Title')
    doc.add_paragraph('DeepSeek问答结果汇编  2026年10月15日',style='Subtitle')
    doc.add_paragraph('本文收录原问题库的全部160题及当前问答系统逐题返回的回答，按原有10个类别编排，供阅读、核对和试用。回答保留原题、资料引用与定位；资料不足、条件不明或生成未通过检查的题目如实列出。')
    counts=report['counts'];errors=counts.get('model_error',0);gaps=counts.get('no_evidence',0);answers=counts.get('llm',0)
    doc.add_paragraph(f'本次运行有{answers}题返回带资料引用的生成回答，{gaps}题返回资料或条件不足提示，{errors}题未取得通过检查的生成回答。该统计描述系统运行表现，不作为逐题正确性或完整性验收结论。')
    doc.add_paragraph('资料编号在每道题内独立使用，对应本题末尾列出的来源。历史材料只适用于其记载的时间和口径；涉及当前预警或业务签发时，应核对官方实际发布。')
    p=doc.add_paragraph();inline(p,'生成时间  '+report['completed_at'][:19].replace('T','  ')+'  北京时间',10)
    p=doc.add_paragraph();inline(p,'问答模型  '+report['model']+'    原题来源  '+report['source_name'],10)
    doc.add_paragraph('分类索引',style='Heading 2')
    records=[['类别','题数']]
    for name,items in groups.items(): records.append([name,str(len(items))])
    table(doc,records,[5.8,.7])
    for index,(category,items) in enumerate(groups.items(),1):
        if index==1: doc.add_page_break()
        else: doc.styles['Heading 1'].paragraph_format.space_before=Pt(18)
        doc.add_paragraph(f'{index:02d}  {category}',style='Heading 1')
        p=doc.add_paragraph();p.paragraph_format.keep_with_next=True
        inline(p,f"本类共{len(items)}题  问题编号沿用工作台原编号",10)
        for item in items:
            paragraph_start=len(doc.paragraphs)
            number=item['id'].replace('-','')
            sub=re.sub(r'[^\w\u4e00-\u9fff ]',' ',item.get('subrequirement','')).strip()
            doc.add_paragraph(number+'  '+sub,style='Heading 2')
            p=doc.add_paragraph();p.paragraph_format.keep_with_next=True
            run=p.add_run('问题  ');set_font(run,bold=True);inline(p,item['question'])
            result=item['result'];meta=result.get('answer_meta',{})
            mode=meta.get('used_mode');coverage=meta.get('answerability')
            label='生成回答' if mode=='llm' else '资料或条件不足' if mode=='no_evidence' else '本次生成未完成'
            if coverage=='partial': label+='  部分子问题尚缺依据'
            p=doc.add_paragraph();p.paragraph_format.keep_with_next=True;inline(p,'系统回答  '+label,10)
            answer=result.get('answer') or '本次没有取得可展示的系统回答。'
            add_answer(doc,answer)
            sources=sorted(result.get('rag_sources',[])+result.get('web_sources',[]), key=lambda row:row.get('citation_id',0))
            if sources:
                doc.paragraphs[-1].paragraph_format.keep_with_next=True
                p=doc.add_paragraph();p.paragraph_format.keep_with_next=True;inline(p,'本题资料来源',9.5)
                for source_index,row in enumerate(sources):
                    p=doc.add_paragraph();p.paragraph_format.space_after=Pt(2)
                    p.paragraph_format.keep_together=True
                    p.paragraph_format.keep_with_next=source_index<len(sources)-1
                    inline(p,'['+row['citation']+'] '+str(row.get('title') or row.get('filename') or '来源资料')+'  '+str(row.get('locator') or '定位未注明'),9.5)
                    if row.get('url'):
                        inline(p,'  '+row['url'],9.5)
            question_paragraphs=doc.paragraphs[paragraph_start:]
            if sum(len(p.text) for p in question_paragraphs)<850 and '|' not in answer:
                for p in question_paragraphs[:-1]:
                    p.paragraph_format.keep_with_next=True
                    p.paragraph_format.keep_together=True
                question_paragraphs[-1].paragraph_format.keep_with_next=False
    h=section.header.paragraphs[0];inline(h,'海洋灾害160题问答汇编',9)
    f=section.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.CENTER
    inline(f,'第 ',9);field(f,'PAGE');inline(f,' 页  共 ',9);field(f,'NUMPAGES');inline(f,' 页',9)
    settings=doc.settings._element
    update=OxmlElement('w:updateFields');update.set(qn('w:val'),'true');settings.append(update)
    doc.core_properties.title='海洋灾害160题与系统回答'
    doc.core_properties.subject='当前DeepSeek问答系统逐题回答汇编'
    doc.core_properties.author='海洋灾害问答项目'
    doc.core_properties.keywords='海洋灾害 台风 风暴潮 海浪 160题 DeepSeek'
    DESTINATION.parent.mkdir(parents=True,exist_ok=True)
    doc.save(DESTINATION)
    print(json.dumps({'document':str(DESTINATION),'questions':160,'categories':len(groups),'counts':counts},ensure_ascii=False))


if __name__=='__main__': main()

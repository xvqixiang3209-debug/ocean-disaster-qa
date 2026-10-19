"""Word report of checkpointed actual replies, coverage, sources and unresolved inputs."""
import json,re
from collections import Counter,OrderedDict
from pathlib import Path
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from build_160_qa_document import set_font,inline,field,table,add_answer

APP=Path(__file__).resolve().parents[1]
RUN=APP/'tmp/qa_eval_20261017'
DEST=APP/'exports/海洋灾害160题实测问答与数据缺项_20261017.docx'
LABELS=['基本可答','假设试算可答','部分可答','依据不足','生成失败']

def assess(item):
    r=item['result'];m=r.get('answer_meta',{});mode=m.get('used_mode');coverage=m.get('answerability')
    label=('生成失败' if mode=='model_error' else '部分可答' if coverage=='partial' else
           '假设试算可答' if coverage=='full' and m.get('business_trial') else
           '基本可答' if mode=='llm' and coverage=='full' else '依据不足')
    parts=m.get('subquestion_coverage') or []
    if label=='部分可答' and parts and all(p.get('status')=='insufficient' for p in parts):label='依据不足'
    if m.get('business_trial') and parts and all(p.get('status')=='full' for p in parts) and not m.get('missing_information'):
        label='假设试算可答'
    gaps=list(m.get('missing_information') or [])
    for part in m.get('subquestion_coverage',[]):
        for gap in part.get('missing_information') or []:
            if gap and gap not in gaps:gaps.append(gap)
    notes=[]
    if m.get('business_trial'):notes.append('业务数据按用户授权假设试算，单位、高程基准、站点映射和时间口径尚未全部核定。')
    if label=='生成失败':
        gaps=['本次未取得通过检查的系统回答，不能仅凭失败判断资料是否缺失。']
        notes.append('失败原因：'+str(m.get('validation_error') or m.get('error_type') or m.get('route') or '未注明'))
    if label=='依据不足' and not gaps:gaps=['系统未取得足够依据，或具体地点、时段、授权信息不足；见原答。']
    if m.get('web_supplement_state')=='unavailable':notes.append('本题联网检索未形成可用证据，不等于互联网上不存在相关资料。')
    if coverage=='partial' and label=='依据不足':notes.append('系统总标记为部分覆盖，但各子问题均标记依据不足，本报告按子问题归为依据不足。')
    if m.get('business_trial') and label=='假设试算可答' and coverage!='full':notes.append('系统总标记为部分覆盖，但本题各子问题均有完整返回，本报告列为假设口径下可答。')
    sources=r.get('rag_sources',[])+r.get('web_sources',[])
    if any(s.get('review_label')=='历史资料，未审定' for s in sources):notes.append('本题含历史归档资料，不能据此确认当前预警、现行版本或本次灾情。')
    if item['id'] in {'Q-051','Q-052','Q-053','Q-054'}:notes.append('系统自身能力、更新频率和准确率须结合运行配置与独立评测核验，公开行业资料不能替代。')
    if item.get('review_override'):
        override=item['review_override'];label=override.get('label',label);gaps+=override.get('gaps',[]);notes+=override.get('notes',[])
    return {'label':label,'gaps':list(dict.fromkeys(gaps)),'notes':list(dict.fromkeys(notes))}

def para(doc,text='',size=11,bold=False,keep=False):
    p=doc.add_paragraph();p.paragraph_format.keep_with_next=keep
    inline(p,('**'+text+'**') if bold else text,size);return p

def link(p,url,label='网页原文'):
    h=OxmlElement('w:hyperlink');h.set(qn('r:id'),p.part.relate_to(url,RT.HYPERLINK,is_external=True))
    r=OxmlElement('w:r');pr=OxmlElement('w:rPr');sz=OxmlElement('w:sz');sz.set(qn('w:val'),'19');pr.append(sz)
    color=OxmlElement('w:color');color.set(qn('w:val'),'165C75');pr.append(color);r.append(pr)
    t=OxmlElement('w:t');t.text=label;r.append(t);h.append(r);p._p.append(h)

def main():
    report=json.loads((RUN/'answers.json').read_text(encoding='utf-8'))
    manifest=json.loads((RUN/'manifest.json').read_text(encoding='utf-8'))
    snapshot=json.loads((RUN/'snapshot_start.json').read_text(encoding='utf-8'))
    rows=report['items'];assert len(rows)==160 and len({x['id'] for x in rows})==160
    overrides_path=RUN/'review_overrides.json'
    overrides=json.loads(overrides_path.read_text(encoding='utf-8')) if overrides_path.exists() else {}
    for row in rows:
        row['review_override']=overrides.get(row['id']);row['assessment']=assess(row)
    counts=Counter(x['assessment']['label'] for x in rows)
    groups=OrderedDict()
    for row in rows:groups.setdefault(row['category'],[]).append(row)
    doc=Document();s=doc.sections[0]
    s.page_width=Inches(8.5);s.page_height=Inches(11)
    s.left_margin=s.right_margin=Inches(1);s.top_margin=s.bottom_margin=Inches(.8)
    s.header_distance=s.footer_distance=Inches(.35)
    for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','List Bullet','List Number']:
        st=doc.styles[name];st.font.name='Calibri';st.font.color.rgb=RGBColor(0,0,0);st.font.italic=False
        st._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'宋体')
        pp=st._element.get_or_add_pPr()
        for e in list(pp.findall(qn('w:pBdr'))):pp.remove(e)
    n=doc.styles['Normal'];n.font.size=Pt(11);n.paragraph_format.line_spacing=1.15;n.paragraph_format.space_after=Pt(5);n.paragraph_format.widow_control=True
    for name,size,before,after in [('Title',22,0,12),('Subtitle',11,0,12),('Heading 1',16,16,8),('Heading 2',12,12,6)]:
        st=doc.styles[name];st.font.size=Pt(size);st.font.bold=name!='Subtitle'
        st.paragraph_format.space_before=Pt(before);st.paragraph_format.space_after=Pt(after);st.paragraph_format.keep_with_next=True
    doc.add_paragraph('海洋灾害160题实测问答与数据缺项',style='Title')
    doc.add_paragraph('2026年10月17日  当前系统实际返回结果',style='Subtitle')
    para(doc,f'本次对原问题库全部160题逐题调用当前系统，保留原题、具体化后的实测题、完整原答和引用来源。结果为基本可答{counts["基本可答"]}题、假设试算可答{counts["假设试算可答"]}题、部分可答{counts["部分可答"]}题、依据不足{counts["依据不足"]}题、生成失败{counts["生成失败"]}题。')
    para(doc,'这些结果是当前系统的运行覆盖情况。分类参考系统返回的子问题覆盖标记，并结合题意核对明显缺项；没有把生成成功视为答案正确，也没有完成专家逐条正确性验收。正文保留失败提示和不足之处，不以人工撰写答案替换系统原答。')
    doc.add_paragraph('测试范围与具体化方式',style='Heading 1')
    para(doc,'测试使用本地正在运行的问答服务，调用与界面共用的问答引擎；每题开启联网补充并独立提问，不带上一题对话历史。模型为'+report['model']+'，版本为phase1.4.0-forecast-trial。')
    para(doc,'原题中的“这次台风”主要替换为接口活动台风'+manifest['storm']['name']+'（'+manifest['storm']['sid']+'）；“我市”“XX市XX区”等补充为宁波市、鄞州区，海域以舟山为代表。未来72小时预报固定为'+manifest['forecast_interval']+'，不是把每道题的提问时刻向后滚动。')
    para(doc,'明确标记“假设”“情景”的问题用于测试避险、停运、理赔等流程，不表示这些灾情或命令真实发生。R-026以2019年利奇马浙江防御过程做历史复盘；其他历史题保留统计范围和年份。未把所有当前问题替换成容易回答的历史问题。')
    para(doc,'采集开始 '+report['started_at']+'；采集结束 '+report['completed_at']+'。每道题保留各自的返回时间。资料编号在本题内独立使用，与其他题不共用。',10)
    definitions=[['分类','本报告中的含义'],['基本可答','返回有依据的回答，系统覆盖标记为完整；仍需核验准确性和版本。'],['假设试算可答','业务接口或网格计算能给出结果，但采用尚待核定的假设口径。'],['部分可答','至少有子问题可解释，仍缺资料、实时记录或有效结论。'],['依据不足','没有取得足够证据，不能形成该题所需结论。'],['生成失败','请求、模型生成或引用校验失败；与“缺数据”分开统计。']]
    table(doc,definitions,[1.1,5.4])
    doc.add_paragraph('各类问题覆盖情况',style='Heading 1')
    records=[['类别','题数','基本\n可答','假设\n试算','部分\n可答','依据\n不足','生成\n失败']]
    for cat,items in groups.items():
        c=Counter(x['assessment']['label'] for x in items);records.append([cat,str(len(items))]+[str(c[k]) for k in LABELS])
    records.append(['合计','160']+[str(counts[k]) for k in LABELS]);table(doc,records,[2.3,.55,.73,.73,.73,.73,.73])
    doc.add_paragraph('缺项与后续接入内容',style='Heading 1')
    para(doc,'缺项分为数据尚未接入、资料已存在但未检索到或未进入正确业务路由、现行版本待核以及生成校验失败。联网未取得可用正文不表示官方资料不存在；不能据本次失败要求重新导入全部资料。')
    for title,text in [
      ('当前预警与业务命令','接入带发布时间、有效时段、覆盖海域的正式海浪和风暴潮警报，以及响应、五停、转移、回港、复航公告。活动台风和网格结果不能替代正式发布。'),
      ('站点预报与口径核定','核定SWL变量含义、天文潮高程基准、单位、时区、站点映射、岸段阈值版本和无效等级。验潮站与岸段不能互换，平均风、阵风和风力等级也不能直接替代。'),
      ('风险与现场业务台账','接入淹没区和DEM、海堤达标及缺口、承灾体和人口、避护场所容量及开放状态、停复运和抢修进度等。通用防御方法不能得出某村转移人数或某住户是否进水。'),
      ('观测与预报检验','接入潮位、浮标、雷达等实时观测及缺测状态、集合成员、波向和谱峰周期等受支持变量；提供2023至2025年预警和实况配对记录，计算实际评分与提前量。天文潮和NC预报不算实时观测。'),
      ('统计与政策版本','补齐指定年份灾损公报、现行地方预案和标准有效性信息、保险条款与救助办理渠道。归档官方材料可以解释历史条款，不能自动视为现行规定。'),
      ('检索与生成问题','对有资料但回答仍不完整的题检查多子问题召回和业务路由；对生成失败题检查引用、数字和覆盖校验，不能将这些失败直接归因为没有数据。')]:
        para(doc,title,bold=True,keep=True);para(doc,text)
    doc.add_paragraph('本次发现的程序问题',style='Heading 1')
    para(doc,'Q-012明确要求跨日72小时时段，程序实际查询变成10月17日00时至23时。Q-022同样发生时段缩短，并返回198个岸段全部缺数；起始快照已有覆盖未来多天的NC文件，因此不能把这些结果直接理解为未下载预报数据。')
    para(doc,'Q-051、Q-052询问系统能力和数据更新，却进入活动台风查询路由，回答只覆盖台风接口。这是意图路由和系统说明未接入的问题。Q-017明确给出“假设发布红色警报”，回答仍要求验证当前是否发布，说明假设情景与实时查询的识别需要区分。')
    failure_counts=Counter(x['result'].get('answer_meta',{}).get('validation_error') or x['result'].get('answer_meta',{}).get('error_type') or x['result'].get('answer_meta',{}).get('route','未注明') for x in rows if x['assessment']['label']=='生成失败')
    para(doc,'生成失败中各原因的返回次数为：'+'；'.join(str(k)+' '+str(v)+'题' for k,v in failure_counts.items())+'。其中unsupported_insufficient_answer表示依据不足类回答触发校验拒绝，不表示模型服务没有连接。')
    para(doc,'Q-024引用了程序汇总的“历史总结-2026-BAVI.txt”。此类汇总须回查原报告的台风身份、年份、实测与预报字段，不能只凭程序生成的标题确认真实历史事件。该题原答完整保留，未将其中事件身份视为已核验事实。')
    doc.add_paragraph('数据快照与试算限制',style='Heading 1')
    f=snapshot['forecast'];imp=f.get('import',{})
    para(doc,f'本次起始快照有{imp.get("sites")}个站点登记、{imp.get("coasts")}个岸段、{imp.get("tide_files")}份天文潮文件、{imp.get("tide_hours")}条逐时天文潮记录和{f.get("grid_files")}份已登记网格文件。网格产品包括SWL、WAV、SSW、OCU、SST；数量是文件登记数，不是可回答题目数。')
    para(doc,'试算沿用用户已授权假设：SWL按纯增水并采用米，天文潮原值按已订正至1985基准使用，四色岸段阈值按厘米和1985基准解释，无时区时间按北京时间处理。缺失阈值和无效值不补造。所有这些数值结论均保留假设口径标记。')
    para(doc,'同步快照在10月17日15时已列出文件名包含当日20时的L2起报批次。文件名时区及起报时间约定需要核定，不能仅凭“最新文件”确认它对应当前有效预报。报告保留各题实际批次及时间，不掩盖这一时间口径问题。')
    para(doc,'下文引用为系统回答实际使用的来源。来源中的历史、未审定、版本未核或解析缺项提示均影响可用范围。任何试算或一般方法解释都不构成正式预警、行政决定或现场安全许可。')
    doc.add_paragraph('逐题结果索引',style='Heading 1')
    for label in LABELS:
        ids=[x['id'] for x in rows if x['assessment']['label']==label]
        para(doc,f'{label}  {len(ids)}题',bold=True,keep=True);para(doc,'、'.join(ids) if ids else '无',10)
    doc.add_page_break()
    for i,(cat,items) in enumerate(groups.items(),1):
        doc.add_paragraph(f'{i:02d} {cat}',style='Heading 1')
        for item in items:
            a=item['assessment'];r=item['result'];m=r.get('answer_meta',{})
            doc.add_paragraph(item['id']+' '+re.sub(r'[^\w\u4e00-\u9fff ]',' ',item.get('subrequirement','')).strip(),style='Heading 2')
            if item['question']!=item['original_question']:
                para(doc,'原题  '+item['original_question'],10,keep=True)
                para(doc,'实测问题  '+item['question'],11,True,True)
            else:para(doc,'原题及实测问题  '+item['question'],11,True,True)
            para(doc,'测试结果  '+a['label']+'  |  '+item['collected_at']+'  |  耗时'+str(item['elapsed_seconds'])+'秒',9.5,keep=True)
            if a['gaps']:
                para(doc,'尚缺数据或依据',10,True,True)
                for gap in a['gaps']:para(doc,'• '+gap,10)
            if a['notes']:para(doc,'核对说明  '+' '.join(a['notes']),9.5)
            para(doc,'系统原答',10,True,True)
            add_answer(doc,r.get('answer') or '本次没有取得系统回答。')
            sources=sorted(r.get('rag_sources',[])+r.get('web_sources',[]),key=lambda x:x.get('citation_id',0))
            if sources:
                para(doc,'本题引用来源',9.5,True,True)
                for src in sources:
                    title=src.get('title') or src.get('filename') or '未命名来源'
                    text='['+str(src.get('citation') or '资料'+str(src.get('citation_id','')))+'] '+title
                    locator=src.get('locator');text+=('  '+str(locator)) if locator else ''
                    if src.get('review_label'):text+='  '+src['review_label']
                    p=para(doc,text,9.5);p.paragraph_format.space_after=Pt(2)
                    if src.get('url'):inline(p,'  ',9.5);link(p,src['url'])
                    elif src.get('source_id') and not src['source_id'].startswith('forecast:'):
                        # Preserve source locator; local server links are not stable off this machine.
                        pass
            if m.get('business_evidence_id'):para(doc,'业务查询记录号  '+m['business_evidence_id']+'；完整查询快照已随本次实测原始记录留存。',9.5)
    h=s.header.paragraphs[0];inline(h,'海洋灾害160题实测问答与数据缺项',9)
    footer=s.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    inline(footer,'第 ',9);field(footer,'PAGE');inline(footer,' 页  共 ',9);field(footer,'NUMPAGES');inline(footer,' 页',9)
    upd=OxmlElement('w:updateFields');upd.set(qn('w:val'),'true');doc.settings._element.append(upd)
    doc.core_properties.title='海洋灾害160题实测问答与数据缺项';doc.core_properties.author='海洋灾害问答项目'
    DEST.parent.mkdir(parents=True,exist_ok=True);doc.save(DEST)
    evaluated={**report,'assessment_counts':dict(counts),'items':rows}
    (RUN/'evaluated_answers.json').write_text(json.dumps(evaluated,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    print(json.dumps({'output':str(DEST),'counts':dict(counts),'questions':160},ensure_ascii=False))

if __name__=='__main__':main()

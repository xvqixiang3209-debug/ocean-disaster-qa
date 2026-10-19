"""Create the scope/validity evaluation report from immutable actual answers."""
from collections import Counter
from datetime import datetime
import hashlib,json
from pathlib import Path
import sys

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'tools'))
from build_160_phase16_document import (configure,paragraph,labeled,table,bookmark,
    markdown_answer as _markdown_answer,source_paragraphs,review_items)
from docx import Document

RUN=ROOT/'tmp/qa_full160_phase162_20261018_final'
OUTPUT=ROOT/'exports/联网检索改造与160题前后对比及实际回答_20261018.docx'
LABELS=('完整回答','部分回答','未回答')


def load(path):return json.loads(path.read_text(encoding='utf8'))


def markdown_answer(doc,answer):
    return _markdown_answer(doc,answer,wide_as_table=True)


def scope_note(doc,source):
    app=(source.get('metadata') or {}).get('applicability') or {}
    if not app:return
    use={'direct':'候选直接依据','reference':'限定参考','excluded':'排除'}.get(app.get('use'),app.get('use'))
    labeled(doc,'适用性记录',str(use)+'；地点='+str(app.get('geography'))+'；对象='+str(app.get('object_match'))+
            '；时间='+str(app.get('temporal'))+'；'+('；'.join(app.get('reasons',[])) or '未列出额外限制。'))


def main():
    data=load(RUN/'reviewed_answers.json');summary=data['summary'];rows=data['items']
    transfer=load(RUN/'transfer_answers.json')['items']
    doc=Document();configure(doc)
    doc.add_paragraph('海洋灾害问答联网检索改造与复测报告',style='Title')
    doc.add_paragraph('2026年10月18日  地点对象匹配  资料有效时间  子问题证据覆盖',style='Subtitle')
    before=summary['comparison_counts']['before'];after=summary['comparison_counts']['after']
    paragraph(doc,'本次改造面向问题类型：概念解释、专业方法、预案程序、当前预报预警、场所开放、船舶航线及多个子问题。原160题用于回归，另外20条真实接口问题变换地点、日期和问法，自动回归用例还包括未预置的浴场名称。运行代码没有按题号匹配或保存固定答案。')
    paragraph(doc,f'同一评阅口径下，改前完整{before[LABELS[0]]}题、部分{before[LABELS[1]]}题、未回答{before[LABELS[2]]}题；改后完整{after[LABELS[0]]}题、部分{after[LABELS[1]]}题、未回答{after[LABELS[2]]}题。该结果为辅助AI对实际回答覆盖程度的盲评，不是专业正确率或业务专家验收。')
    paragraph(doc,'本次处理了证据的适用边界，没有补齐所有未接入的数据。正式公告缺失、预报口径未核、现场开放及个人业务条件不足的题，仍可能只有部分回答。返回“未知”与提供查询方法，不作为已完整交付具体数据或状态。')
    doc.add_paragraph('改造内容与适用的问题类型',style='Heading 1')
    table(doc,[['能力','实现内容','适用的问题类型'],
        ['地点和对象','地区包含关系、具体浴场码头航线名称、浴场与岸钓船钓、客运船与休闲渔船分别匹配；未知名称保留原名核对。','场所状态、停复航、地方规则、活动建议'],
        ['资料时间','发布时间、查询截至时刻、起报与有效时段、生效及废止状态分别判断；未知效力保留为参考。','预报预警、公告、预案标准、历史查询'],
        ['后续公告','仅在同一已识别对象和地区内检查较后生效公告；其他场所的开放通知不能取消本场所的关闭通知。','浴场开放、停航复航、封闭解除'],
        ['子问题覆盖','各子问分别检索与标注候选证据；无直接候选时进行一次有界补查；排序保留不同子问的来源。','条件加决定主体加公告、多个活动、跨地区比较'],
        ['生成与校验','DeepSeek收到来源适用性和候选覆盖表；仅有参考证据的子项不能直接声明完整支持。','综合回答和缺项说明'],
        ['基础知识','保留权威定义页读取；定义来源在概念比较中优先，地名不会阻断通用科普。','定义、成因、区别和常识']], [1.05,3.9,1.65])
    paragraph(doc,'行政区和部分已接入场所采用可扩展目录；未知行政区名称、明确浴场名及航线名采用通用提取。模糊简称、未登记范围和未提供坐标仍可能需要澄清，不能据此声称已覆盖浙江每一个场所。来源正文读取失败、仅有搜索摘要、网页没有有效时段或现行效力信息时，保留相应限制。')
    doc.add_paragraph('测试方式与比较边界',style='Heading 1')
    paragraph(doc,'自动回归结果：539项Python测试及102个子测试通过。其中51个新增参数化测试覆盖日期、地区、场所、业务对象、版本、子问边界与适用性信息进入模型的完整流程。测试记录与本轮真实接口响应分别保存；程序检查不替代业务验收。')
    paragraph(doc,'160题通过与网页相同的 /api/assistant/ask 接口逐题独立调用，沿用改前具体化问题，没有删除未答子问。改前使用已保存的phase1.6.1响应，改后使用phase1.6.2响应。预报后台自动下载在测试时关闭，本轮不验证自动更新时效。运行指纹覆盖14个文件，测试后再核对文件内容未改变。')
    paragraph(doc,'完整性采用同一DeepSeek模型和同一评分规则，逐题将两份回答随机放入A/B位置复核，不向评阅模型提供版本新旧及历史评分；每题使用相同核心要求，结果绑定实际回答SHA256。模型也可能误评，报告保留原答和证据限制；没有开展业务专家逐题验收。')
    paragraph(doc,'这是一次真实环境的前后对比。外部搜索、网站正文和网络可用性可能变化，答案生成也有波动，因此分类变化不能全部归因于代码，也不能解释为统计显著的准确率提升。')
    replay=load(RUN/'scope_replay.json')
    paragraph(doc,f'另以相同人工构造的搜索候选和正文回放改前、改后检索模块，{replay["passed"]}/{replay["case_count"]}组达到预期。此项隔离网页和网络波动，检查候选保留及适用性标注，不评估模型最终答案或真实网站可用性。')
    table(doc,[['控制场景','改前候选','改后处理']]+[[r['family'],
        '保留但未标适用性' if r['before_candidate_retained'] else '未保留',
        {'excluded':'排除','reference':'仅参考','direct':'可作候选直接依据'}.get(r['after_use'],r['after_use'])]
        for r in replay['items']], [2.5,1.9,2.2])
    table(doc,[['分类','改前重新盲评','改后盲评']]+[[k,str(before[k]),str(after[k])] for k in LABELS],[3.4,1.6,1.6],center_cols=(1,2))
    original=summary['original_previous_review_counts']
    paragraph(doc,'上一份报告原有分类为'+str(original)+'。本报告为保证口径一致，重新盲评两份实际回答，以上对比不与上一份报告的原分类混算。')
    metrics=summary['observed_metrics']
    table(doc,[['可直接观测的指标','改前','改后']]+[
        [label,str(metrics['before'][key]),str(metrics['after'][key])] for label,key in (
            ('HTTP成功数','http_200'),('使用联网来源的题数','questions_using_web'),
            ('返回联网来源条目数','returned_web_sources'),('平均响应秒数','mean_seconds'))], [3.4,1.6,1.6],center_cols=(1,2))
    paragraph(doc,'联网来源数量仅反映取得并引用的条目数，不直接等于相关性或正确性。改后来源保留程序适用性标记：'+str(metrics['after']['web_applicability_labels'])+'；时间标记：'+str(metrics['after']['web_temporal_labels'])+'。标为reference的来源仍可能用于解释历史或一般条件，不能据此证明当前开放或规则现行。')
    paragraph(doc,'本轮盲评偏好统计：改后较好'+str(summary['pair_preference'].get('after',0))+'题，改前较好'+str(summary['pair_preference'].get('before',0))+'题，持平'+str(summary['pair_preference'].get('tie',0))+'题。该主观偏好不代表正确率，完整性总量本轮没有提高。')
    computation=load(RUN/'computation_audit.json')
    paragraph(doc,'程序结果检查：11题必需事实交付通过，预报区间不一致及引用编号缺失均未发现；语言说明核对失败后返回程序表的题号为'+ '、'.join(computation['computed_table_fallback_ids'])+'。首次完整测试Q-040发生引用数量校验失败，该次HTTP503及空正文保留在逐题记录中。')
    doc.add_paragraph('按问题类型查看覆盖情况',style='Heading 1')
    category_rows=[['原问题类别','改前完整 部分 未答','改后完整 部分 未答']]
    for category,counts in summary['category_comparison'].items():
        category_rows.append([category,' / '.join(str(counts['before'].get(k,0)) for k in LABELS),
                              ' / '.join(str(counts['after'].get(k,0)) for k in LABELS)])
    table(doc,category_rows,[3.2,1.7,1.7],center_cols=(1,2))
    paragraph(doc,'同类问题还需要相同的数据链路：基础概念依赖可读取的权威定义；预案程序依赖适用地区和现行条款；数值题依赖NC及确定的变量单位和基准；开放与停复航依赖对应场所或航线的正式通知；转移人数、损失及淹没范围依赖额外的业务台账与计算模块。检索改造提高证据选择能力，不会自动产生缺失数据。')
    doc.add_paragraph('补充科普资料后的追加验证',style='Heading 1')
    post=load(RUN/'post_supplement/results.json')['items']
    added=load(RUN/'public_concept_supplement/imported.json')
    new_ids={row['source_id'] for row in added}
    paragraph(doc,'正式160题测试结束后，补入两篇经公开页面核对的科普摘要，分别涉及台风季节和重现期。摘要明确标注非逐字原文，保留来源网址、原网页发布时间和入库时间，仅用于通用知识，不作为地方行动命令或现行工程标准。未在160题运行中改变资料库，也未用追加结果替换失败原答或重算总分。')
    for row in added:paragraph(doc,row['filename']+'；来源：'+row['url']+'；入库：'+row['imported_at'])
    post_table=[['追加题号','原提问或变体','HTTP','命中新增资料','接口覆盖标签']]
    for row in post:
        hit=any(source.get('source_id') in new_ids for source in row['reply'].get('sources',[]))
        post_table.append([row['id'],row['question'],str(row['http_status']),'是' if hit else '否',str(row['reply'].get('answerability'))])
    table(doc,post_table,[.65,3.2,.45,1.1,1.2],center_cols=(0,2,3,4))
    paragraph(doc,'追加观察：Q-040已实际返回带来源的解释，另两个重现期变体P-003、P-004也命中新增资料并返回正文。台风季节原题Q-003及福建变体P-001仍未命中新资料，浙江变体P-002能够命中。这说明资料入库后仍存在同义问法召回和通用气候知识排序问题；没有把“入库成功”当作“同类问题全部解决”。表内full/partial仅为接口自报字段，追加六题未采用160题的成对盲评。')
    doc.add_paragraph('发生分类变化的问题',style='Heading 1')
    changes=summary['changed']
    table(doc,[['题号','改前','改后','变化理由']]+[[v['id'],v['before'],v['after'],v['reason']] for v in changes], [.65,.85,.85,4.25])
    doc.add_paragraph('全部160题逐题对照',style='Heading 1')
    table(doc,[['题号','问题类别','改前','改后']]+[[r['id'],r['category'],r['assessment']['previous_label'],r['assessment']['label']] for r in rows], [.7,3.7,1.1,1.1],center_cols=(0,2,3))
    doc.add_paragraph('160题系统实际回答',style='Heading 1')
    audits=[]
    for number,row in enumerate(rows,1):
        review=row['assessment'];reply=row['reply']
        p=doc.add_paragraph(row['id']+' '+row.get('subrequirement',''),style='Heading 2')
        bookmark(p,'q_'+row['id'].replace('-','_'),number)
        labeled(doc,'实际提问',row['question'])
        labeled(doc,'完整性对比',review['previous_label']+' → '+review['label'])
        labeled(doc,'复核理由',review['reason'])
        review_items(doc,'尚缺的核心输出或依据',review['missing_points'])
        if review['evidence_checks']:
            review_items(doc,'证据限制',review['evidence_checks'])
        labeled(doc,'调用记录',f"{row['started_at']}；HTTP {row['http_status']}；模式 {reply.get('mode')}；试算 {reply.get('trial')}。")
        answer=reply.get('answer','')
        if not answer:labeled(doc,'接口未返回正文',reply.get('error') or '本次接口未返回answer字段。')
        paragraph(doc,'系统实际原答',bold=True,keep_next=True)
        if answer:audit=markdown_answer(doc,answer)
        else:
            audit={'raw_sha256':hashlib.sha256(b'').hexdigest(),'preservation_check':'pass'}
        sources=reply.get('sources',[])
        paragraph(doc,f'本题实际来源  {len(sources)}条',bold=True,keep_next=True)
        for i,source in enumerate(sources,1):
            source_paragraphs(doc,source,i)
            scope_note(doc,source)
        if not sources:paragraph(doc,'接口未返回来源条目。')
        labeled(doc,'追溯记录',str(RUN/(row['id']+'.json'))+'；SHA256 '+audit['raw_sha256'])
        audits.append({'id':row['id'],**audit})
    doc.add_paragraph('另20条真实变体问题',style='Heading 1')
    paragraph(doc,'以下不属于原160题，不计入前后完整性数量。每条均改变日期、地点、对象组合或问法，保存真实返回正文；这20条没有采用与160题相同的成对盲评，不能据此声称新增20条完整正确回答。')
    paragraph(doc,'变体观察：T-006把浙江归档预案、东海局程序和缺少福建省依据分别说明，没有把上级机构流程当作福建地方预案。T-009区分景区营业时间与浴场游泳开放。T-008没有用其他月份预警代替历史截至时点的预警。以上表现支持边界表达改进，不说明缺失的公告或规则已补齐。')
    paragraph(doc,'变体仍暴露上游能力限制：T-011同时提到南沙与东沙时，业务定位要求分别查询，未能在一轮中计算南沙并解释东沙消息的适用性；T-015的连云港浴场超出现有浙江定位与预报接入范围；部分活动题仍可能只返回计算表。版本效力题也可能检索到同年份灾情材料而非预案正文。这些问题保留为下一轮业务路由、对象拆分、来源补齐和语言生成改进项，不能宣称同类问题均已完整回答。')
    for row in transfer:
        doc.add_paragraph(row['id']+' '+row['family'],style='Heading 2')
        labeled(doc,'变体提问',row['question'])
        labeled(doc,'实际调用',f"HTTP {row['http_status']}；模式 {row['reply'].get('mode')}；接口自报覆盖 {row['reply'].get('answerability')}，该字段不是独立验收。")
        markdown_answer(doc,row['reply'].get('answer') or row['reply'].get('error','接口未返回正文。'))
        paragraph(doc,'变体来源',bold=True,keep_next=True)
        for i,source in enumerate(row['reply'].get('sources',[]),1):
            source_paragraphs(doc,source,i)
            scope_note(doc,source)
        labeled(doc,'变体追溯',str(RUN/(row['id']+'.json')))
    doc.core_properties.title='海洋灾害问答联网检索改造与复测报告'
    doc.core_properties.author='海洋灾害问答项目'
    OUTPUT.parent.mkdir(exist_ok=True);doc.save(OUTPUT)
    (RUN/'document_audit.json').write_text(json.dumps({'output':str(OUTPUT),
        'output_sha256':hashlib.sha256(OUTPUT.read_bytes()).hexdigest(),'items':audits,
        'all_answers_preserved':all(a['preservation_check']=='pass' for a in audits),
        'visual_review':'pending'},ensure_ascii=False,indent=2),encoding='utf8')
    print(str(OUTPUT))


if __name__=='__main__':main()

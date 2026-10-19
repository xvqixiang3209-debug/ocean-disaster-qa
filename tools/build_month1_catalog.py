"""Build the shipped, unlabelled 160-question requirements catalogue."""
from __future__ import annotations

import argparse
from datetime import date, datetime
import hashlib
import json
from pathlib import Path

from openpyxl import load_workbook


APP = Path(__file__).resolve().parents[1]
DEFAULT_SOURCE = Path(r"E:\ocean_qa\implementation_plan_20261007_business\问题需求分类与三个月验收台账.xlsx")
DEFAULT_OUTPUT = APP / "data" / "requirements" / "catalog.json"


def clean(value):
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return str(value).strip()


def build(source: Path, output: Path) -> dict:
    source = source.resolve(strict=True)
    output = output.resolve()
    workbook = load_workbook(source, read_only=True, data_only=True, keep_links=False)
    try:
        categories_sheet = workbook["10类业务需求总览"]
        cat_rows = categories_sheet.iter_rows(values_only=True)
        cat_headers = [clean(v) for v in next(cat_rows)]
        category_records = []
        for values in cat_rows:
            row = {name: clean(values[index]) if index < len(values) else "" for index, name in enumerate(cat_headers) if name}
            if row.get("主类ID"):
                category_records.append({
                    "id": row.get("主类ID", ""), "name": row.get("业务需求主类", ""),
                    "scope": row.get("需求范围", ""), "visible_result": row.get("看得见的结果", ""),
                    "acceptance_boundary": row.get("验收边界", ""),
                    "review_status": row.get("评审状态", "待业务确认"),
                })
        questions_sheet = workbook["160题需求与验收台账"]
        rows = questions_sheet.iter_rows(values_only=True)
        headers = [clean(v) for v in next(rows)]
        records = []
        for row_number, values in enumerate(rows, 2):
            fields = {name: clean(values[index]) if index < len(values) else "" for index, name in enumerate(headers) if name}
            identity = fields.get("原问题ID")
            question = fields.get("原问题全文")
            if not identity or not question:
                continue
            records.append({
                "id": identity,
                "question": question,
                "category_id": fields.get("主类ID（建议）", ""),
                "category": fields.get("业务主类（建议）", ""),
                "subrequirement": fields.get("子需求（建议）", ""),
                "user_view": fields.get("用户视角（显示）", ""),
                "role": fields.get("原适用角色", ""),
                "priority": fields.get("原优先级", ""),
                "timeliness": fields.get("原时效要求", ""),
                "agent_task": fields.get("智能体需要做什么（建议）", ""),
                "visible_result": fields.get("用户看得见的结果（建议）", ""),
                "dependencies": fields.get("所需资料与交付依赖（建议）", ""),
                "first_acceptance_window": fields.get("首个成果验收窗口（建议）", ""),
                "acceptance_criteria": fields.get("逐题验收要求（建议）", ""),
                "original_workbook_sheet": fields.get("原工作表", ""),
                "original_excel_row": fields.get("原Excel行号", ""),
                "original_note_unverified": fields.get("原备注（未经核实）", ""),
                "linked_requirements": fields.get("关联需求/复合子问题（建议）", ""),
                "status": "pending_confirmation",
                "status_label": "待业务核定",
                "status_note": "原台账未提供经业务确认的逐题验收状态；请先核定题目、资料和计分口径。",
            })
    finally:
        workbook.close()
    if len(records) != 160 or len(category_records) != 10:
        raise ValueError(f"预期160题和10类，实际为{len(records)}题/{len(category_records)}类")
    if len({item["id"] for item in records}) != len(records):
        raise ValueError("原问题ID重复")
    payload = {
        "schema_version": 1,
        "title": "台风与海洋灾害专业问答需求台账",
        "source_name": "问题需求分类与三个月验收台账.xlsx",
        "source_sha256": hashlib.sha256(source.read_bytes()).hexdigest(),
        "generated_at": datetime.now().astimezone().isoformat(timespec="seconds"),
        "question_count": len(records),
        "category_count": len(category_records),
        "evaluation_passed_count": 0,
        "warning": "需求台账不是知识事实、专家答案或训练数据；未经过业务确认的题目状态均为待业务核定。",
        "categories": category_records,
        "items": records,
    }
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    summary = {"output": str(output), "questions": len(records), "categories": len(category_records),
               "month_one_questions": sum(item["first_acceptance_window"] == "第1个月" for item in records),
               "source_sha256": payload["source_sha256"], "passed_count": 0}
    print(json.dumps(summary, ensure_ascii=False, indent=2))
    return summary


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, default=DEFAULT_SOURCE)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()
    build(args.source, args.output)


if __name__ == "__main__":
    main()

"""Export customer questions as UNLABELLED evaluation seeds, never RAG facts.

No model calls, fabricated reference answers, automatic passing or database
writes. The source workbook is read-only. Candidate routes are transparent
keyword/time-field rules, not validated labels.
"""
from __future__ import annotations

import argparse
from collections import Counter
from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path
import re

from openpyxl import load_workbook

APP = Path(__file__).resolve().parents[1]
DEFAULT_SOURCE = Path(r"E:\ocean_qa\台风海洋灾害智能体基础问题库.xlsx")
DEFAULT_OUTPUT = APP / "docs" / "基础问题库_验收待标注.jsonl"


def clean(value):
    return "" if value is None else str(value).strip()


def candidate_route(question, fields):
    """Suggest required tools without implying data exists or can be answered."""
    timeliness = clean(fields.get("时效要求"))
    context = question + " " + " ".join(clean(fields.get(key)) for key in ("二级分类", "所需数据或知识源"))
    routes, needs = [], []
    if timeliness in {"实时", "准实时"} or re.search(r"今天|明天|当前|正在|这次|本次|最新|未来\d", question):
        routes.append("实时业务数据/官方产品查询")
        needs.append("有授权接口及明确起报、有效时次；未接入时应说明无法核实当前状态")
    if re.search(r"标准|规范|预案|警戒|响应|签发|法律|劳动法|五停|退改|政策|保险|赔付|赔偿", context):
        routes.append("已审核条款RAG/适用性核验")
        needs.append("核验地域、版本、生效时间、触发条件、主体权限；题目备注不是权威条款")
    if re.search(r"多少|几个|次数|概率|超警|排名|平均|最大|最小|最高潮|统计|离散|误差|重现期|面积|经济损失", context):
        routes.append("结构化查询/程序计算")
        needs.append("先确认单位、基准、统计范围和缺测值，再由查询或程序计算，不让模型猜数")
    if re.search(r"淹没|避难|撤离|转移路线|哪里最安全|敏感区|风险区划|空间|海拔|低洼|我家", context):
        routes.append("空间工具/位置与暴露信息")
        needs.append("需要位置、空间图层和适用范围；没有数据时不作确定性安全判断")
    if re.search(r"历史|案例|过去|往年|相似|哪一年|哪次|公报|最大的一次", context):
        routes.append("历史案例RAG/事件检索")
        needs.append("核对台风身份、年份、地域与事件口径，保留资料来源")
    if timeliness not in {"实时", "准实时", "静态", ""}:
        routes.append("定期公报/版本化资料查询")
        needs.append("按所问年份或周期选择完整资料，不能以过期版本冒充最新")
    if not routes or timeliness == "静态":
        routes.append("专业知识RAG/概念解释")
        needs.append("仅依据已知证据回答，规范性结论仍须经过条款审核")
    if re.search(r"安全|危险|能不能|撤离|救|伤|死|眼区|台风眼|触电|涉水|游泳|出海|开闸|调度", question):
        routes.append("安全边界检查/必要时转人工")
        needs.append("不以历史资料保证当下安全，不代替主管部门下达响应、放行或调度命令")
    return {"候选": list(dict.fromkeys(routes)), "依据": f"原表时效要求：{timeliness or '未注明'}；另依据问题及需求字段关键词匹配", "前置条件": list(dict.fromkeys(needs)), "标注状态": "规则建议，待专家确认"}


def build(source, output):
    source = Path(source).resolve(strict=True)
    output = Path(output).resolve()
    if source == output or output.suffix != ".jsonl":
        raise ValueError("Output must be a separate .jsonl file")
    if output.exists():
        # A user may have begun labelling in place. Never erase those answers
        # merely because they rerun the seed generator.
        for line in output.read_text(encoding="utf-8").splitlines():
            if not line.strip():
                continue
            previous = json.loads(line)
            if previous.get("专家参考答案") or previous.get("专家证据来源") or previous.get("专家预期缺失数据处理") or previous.get("专家核定问题族") or previous.get("系统实际回答") or previous.get("验收人") or previous.get("验收意见") or previous.get("pass") is not None or previous.get("验收状态") not in {None, "待验收"}:
                raise ValueError("Existing JSONL contains human labels/results; choose another --output path instead of overwriting")
    sha = hashlib.sha256(source.read_bytes()).hexdigest()
    book = load_workbook(source, read_only=True, data_only=True, keep_links=False)
    records, sheet_counts, time_counts, priority_counts = [], Counter(), Counter(), Counter()
    missing_headers, blank_rows = [], []
    try:
        for sheet in book:
            rows = sheet.iter_rows(values_only=True)
            headers = [clean(value) for value in next(rows, ())]
            if "问题示例" not in headers:
                missing_headers.append(sheet.title)
                continue
            question_index = headers.index("问题示例")
            for row_number, values in enumerate(rows, 2):
                question = clean(values[question_index]) if len(values) > question_index else ""
                if not question:
                    if any(value is not None for value in values):
                        blank_rows.append({"sheet": sheet.title, "row": row_number})
                    continue
                fields = {header: clean(values[index]) if len(values) > index else "" for index, header in enumerate(headers) if header}
                identifier = fields.get("序号") or f"{sheet.title}:{row_number}"
                records.append({
                    "id": identifier,
                    "sheet": sheet.title,
                    "row": row_number,
                    "原题": question,
                    "可能路由": candidate_route(question, fields),
                    "专家参考答案": "",
                    "专家证据来源": [],
                    "专家预期缺失数据处理": "",
                    "专家核定问题族": "",
                    "数据集划分": "未划分，先按问题族/事件分组再划分开发集与留出集",
                    "系统实际回答": "",
                    "pass": None,
                    "验收状态": "待验收",
                    "验收人": "",
                    "验收日期": "",
                    "验收意见": "",
                    "原始需求字段": fields,
                    "source": {"workbook": str(source), "sha256": sha, "question_column": question_index + 1},
                    "用途限制": "仅验收种子；不是知识事实或参考答案；原表备注/追问要求需核实，不导入RAG。",
                })
                sheet_counts[sheet.title] += 1
                time_counts[fields.get("时效要求") or "未注明"] += 1
                priority_counts[fields.get("优先级") or "未注明"] += 1
    finally:
        book.close()
    if not records:
        raise ValueError("No real questions found")
    identities = [(r["sheet"], r["row"]) for r in records]
    if len(set(identities)) != len(identities):
        raise ValueError("Duplicate row provenance")
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text("".join(json.dumps(row, ensure_ascii=False) + "\n" for row in records), encoding="utf-8")
    summary = {
        "generated_at_utc": datetime.now(timezone.utc).isoformat(), "source": str(source), "source_sha256": sha,
        "question_count": len(records), "sheet_counts": dict(sheet_counts), "timeliness_counts": dict(time_counts),
        "priority_counts": dict(priority_counts), "distinct_ids": len({r["id"] for r in records}),
        "duplicate_exact_question_count": len(records) - len({r["原题"] for r in records}),
        "skipped_sheets_without_question_column": missing_headers, "nonempty_rows_without_question": blank_rows,
        "reference_answers_populated": 0, "passed_count": 0, "evaluation_run": False,
        "requirements_D55_N18_recounted": False,
    }
    output.with_name("基础问题库_验收种子统计.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    table = "\n".join(f"| {name} | {count} |" for name, count in sheet_counts.items())
    note = f"""# 基础问题库验收种子说明

本文件由 `tools/build_question_eval.py` 从原始Excel只读生成，实际读取{len(records)}个非空问题，不调用模型，不写入RAG或问答数据库。

| 原工作表 | 问题数 |
|---|---:|
{table}

原文件：`{source}`  
SHA256：`{sha}`

输出 `基础问题库_验收待标注.jsonl` 每行一个问题，保留工作表、Excel原行号、原题、需求字段和来源哈希。可能路由只是透明规则产生的候选，不是已经验收通过的标签。

全部“专家参考答案”留空、`pass=null`、“验收状态=待验收”。没有自动生成权威答案，没有执行{len(records)}题实测，也没有把客户期望/备注中的说法当作真实规范。使用者须核验原备注，例如预案适用条件、定期核定期限和实时追问能力。

推荐流程：业务专家核定问题族与实际问题意图→补充经核验的答案或正确拒答条件→附原文出处和数据时次→按问题族/事件分组后划分开发与留出集→运行系统收集实际答案→逐项判断来源、数字/单位、范围、时效和安全边界。未接入实时接口的题目可以验收“正确说明缺失并引导权威渠道”，不能编造实时答案。

同义改写、同一事件和相同答案事实不得跨开发/留出集合，避免评测泄漏。不要把该JSONL上传到知识工作台当作资料；也不要将它称为已标注训练集。

本Excel只有问题表。本次确认的是{len(records)}题；以前方案中的55项D类需求、18类N类知识源属于其他需求矩阵，不在本工作簿内，未在这里重新验证或合并。

重建命令（在app目录，新环境已安装openpyxl）：

```powershell
..\\.venv\\Scripts\\python.exe .\\tools\\build_question_eval.py
```

该命令重建未标注种子和统计文件；如果检测到已有人工答案/验收记录，会拒绝覆盖。开始人工标注前仍建议另存工作副本。
"""
    output.with_name("基础问题库_验收种子说明.md").write_text(note, encoding="utf-8")
    print(json.dumps(summary, ensure_ascii=False, indent=2))
    return summary


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, default=DEFAULT_SOURCE)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()
    build(args.source, args.output)


if __name__ == "__main__":
    main()

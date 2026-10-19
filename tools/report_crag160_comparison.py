"""Build an immutable comparison dataset, full HTML transcript and Word report."""
from pathlib import Path
from collections import Counter,defaultdict
from datetime import datetime
import argparse,hashlib,html,json,re,sys
ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'tools'))

CHANGES=[
 ('长回答逐句核验','把最多96条既有陈述分为每批24条检查，严格核对每批完整ID；避免活动分析因长度上限退回纯计算表。','marine_kb/claim_grounding.py；marine_kb/qa.py'),
 ('混合提问路由','日期前的“查询”与后续定义子问题分别识别；渠道与模板问法不再误入预警清单。','forecast_runtime/routing.py；forecast_runtime/official_warnings.py'),
 ('海况变量解析','接入MWD原始波向角度，只接受角度单位；同批次同格点同刻才与波高峰值配对。','forecast_runtime/service.py；forecast_runtime/answer_facts.py'),
 ('逐时数据交付','把已有样本数量与覆盖情况送入生成及覆盖检查，检查“有完整序列却说缺序列”的矛盾。','forecast_runtime/service.py；marine_kb/business_validation.py'),
 ('影响报告工具组合','按截止日期选取随后72小时海况，组合台风、岸段和五个沿海地区代表点风浪；区分中心记录与海况预报。','forecast_runtime/service.py'),
 ('历史统计','从历史台风数据库按SID去重生成12个月样本统计，支持年份范围，不解释为浙江登陆频数或概率。','marine_kb/history_statistics.py'),
 ('证据与答案覆盖','未知引用编号进入现有逐句重新归因；业务结果也进行必答点覆盖检查，保留数值和语义校验。','marine_kb/qa.py；marine_kb/coverage_review.py'),
 ('官方来源补充','补充倒灌机理、海平面变化、海域术语、海洋预报服务渠道和2025灾情来源；读取正文后才作为证据。','marine_kb/public_sources.py；marine_kb/web_supplement.py')]
DATA_GAPS=[
 ('省级历史预警完整清单','浙江省接口服务端访问凭据，或按地区与时间归档的发布、更新、解除正文。','国家海洋预报台公开记录已可查询，但不能冒充浙江省级清单或据此判定零预警。'),
 ('预报含义与口径','SWL是否纯增水、SSW语义、潮位高程基准，以及MWD角度零点、旋转方向、来向或去向。','数据有实值，口径尚有假设；角度单位错误文件排除，不能用推测补成确认值。'),
 ('谱峰周期与部分岸段','谱峰周期产品或变量，以及缺失岸段的水位格网或有效关联。','现有NC未发现谱峰周期；部分岸段无有效值，缺测不能按零处理。'),
 ('活动开放和现场条件','指定浴场开放、救生、水质、滩涂出露与返程通道、船舶及经营资质的当天记录。','风浪潮和适用规则可参与筛查，低风浪不能替代实际开放与准入。'),
 ('定量淹没和转移人数','裸地高程、堤防与闸泵、校验后的淹没模型、房屋及人员台账。','地形和人口栅格是输入条件；未运行水动力模型时不能交付淹没深度或正式转移名单。'),
 ('业务检验与损失预估','同站同时观测、模式检验样本、承灾体与经本地验证的损失函数。','历史分项指标与预报峰值可对照，不能据此推出总体灾损轻重或模型命中率。')]

def load(p):return json.loads(p.read_text('utf-8-sig'))
def sha(text):return hashlib.sha256(text.encode('utf-8')).hexdigest()
def save(p,v):p.write_text(json.dumps(v,ensure_ascii=False,indent=2),encoding='utf-8')
def display_blind_text(text,order):
 for tag,name in zip(('A','B'),order):
  text=re.sub(r'(?<![A-Za-z])'+tag+r'(?:\s*版)?(?![A-Za-z])','改造前' if name=='before' else '改造后',str(text))
 return text
def engineering_findings(old,new):
 def grids(value):
  if not isinstance(value,dict):return []
  rows=[value] if value.get('kind')=='grid' else []
  for k in ('items','products'):
   for child in value.get(k,[]) or []:rows+=grids(child)
  return rows
 findings=[];b=old['reply'];a=new['reply'];bg=grids(old.get('business_evidence'));ag=grids(new.get('business_evidence'))
 if not any(g.get('product')=='MWD' and g.get('state')=='ready' for g in bg) and any(g.get('product')=='MWD' and g.get('state')=='ready' for g in ag):
  findings.append({'cause':'工具路由或解析错误','evidence':'同一提问改造前未提取MWD，改造后取得角度单位有效的MWD原始值；现存NC有此变量，原先不能一概归为缺波向数据。','state':'fixed','family':'海况多变量提取'})
 if (b.get('evidence_diagnostics') or {}).get('route')=='official_warnings' and (a.get('evidence_diagnostics') or {}).get('route')!='official_warnings' and re.search(r'官方渠道|查询渠道|哪里|怎么写|模板',new['question']):
  findings.append({'cause':'工具路由或解析错误','evidence':'方法或渠道提问改造前进入预警发布清单，改造后进入资料与网页问答。是否已找全适用正文仍另行判断。','state':'fixed','family':'方法渠道与实时清单区分'})
 if re.search(r'未提供(?:逐时完整|完整逐时|逐时)序列|缺少(?:逐时完整|完整逐时|逐时)序列',b.get('answer','')) and bg and all(g.get('state')=='ready' and g.get('partial') is False and g.get('series') for g in bg):
  findings.append({'cause':'生成遗漏已有证据','evidence':'改造前已保存完整逐时样本，但原答声称未提供逐时序列；属于交付或生成遗漏，不是NC缺测。','state':'fixed' if not re.search(r'未提供逐时|缺少逐时',a.get('answer','')) else 'remaining','family':'逐时覆盖交付'})
 if re.search(r'影响研判报告|影响研判专报|台风影响报告',new['question']) and (old.get('business_evidence') or {}).get('kind')=='typhoons' and (new.get('business_evidence') or {}).get('kind')=='bundle':
  findings.append({'cause':'工具路由或解析错误','evidence':'改造前只调用台风中心接口；改造后同时调用岸段潮位和海域风浪。浙江实况与灾损仍可能缺数，不能把已取得海况预报也写成未提供。','state':'fixed','family':'影响报告工具组合'})
 if (b.get('validation') or {}).get('error')=='missing_or_unknown_citation':
  findings.append({'cause':'请求或校验失败','evidence':'原答保存的校验错误明确为引用编号缺失或越界；新代码对既有陈述逐句重新归因，不能直接据此判定业务资料不存在。','state':'attribution_repair_added','family':'引用编号与逐句证据恢复'})
 return findings
def inline(s):
 s=html.escape(s)
 return re.sub(r'\*\*(.+?)\*\*',r'<strong>\1</strong>',s)
def answer_html(s):
 lines=s.splitlines();out=[];i=0
 while i<len(lines):
  line=lines[i]
  if i+1<len(lines) and '|' in line and re.fullmatch(r'[\s|:\-]+',lines[i+1]) and '-' in lines[i+1]:
   rows=[line];i+=2
   while i<len(lines) and '|' in lines[i] and lines[i].strip():rows.append(lines[i]);i+=1
   out.append('<div class="tablewrap"><table>'+''.join('<tr>'+''.join(f'<{"th" if j==0 else "td"}>{inline(c.strip())}</{"th" if j==0 else "td"}>' for c in row.strip().strip('|').split('|'))+'</tr>' for j,row in enumerate(rows))+'</table></div>');continue
  if line.strip():out.append('<p>'+inline(re.sub(r'^#{1,6}\s+','',line))+'</p>')
  i+=1
 return ''.join(out)
def sources_html(sources):
 out=[]
 for n,s in enumerate(sources,1):
  label=f"[资料{s.get('citation_id',n)}] "+str(s.get('title') or s.get('filename') or '未标标题')+' '+str(s.get('locator') or '')
  url=str(s.get('url') or '')
  if url.startswith('/api/'):url='http://127.0.0.1:8769'+url
  out.append('<li>'+ (f'<a href="{html.escape(url,quote=True)}" rel="noopener noreferrer">{html.escape(label)}</a>' if url.startswith('https://') or url.startswith('http://127.0.0.1:8769/api/') else html.escape(label))+'</li>')
 return '<ol>'+''.join(out)+'</ol>'

def collect(before,after):
 b=load(before/'answers.json');a=load(after/'answers.json');assert len(a['items'])==len(b['items'])==160
 audits={phase:load(folder/'computation_audit_with_final_checker.json') for phase,folder in [('before',before),('after',after)]}
 checks={phase:{r['id']:r for r in audit['items']} for phase,audit in audits.items()}
 selected={phase:load(folder/'selected_passage_audit.json') for phase,folder in [('before',before),('after',after)]}
 assert (before/'manifest.json').read_bytes()==(after/'manifest.json').read_bytes()
 old={r['id']:r for r in b['items']};rows=[]
 for r in a['items']:
  ident=r['id'];o=old[ident];assert r['question']==o['question']
  review=load(after/'paired_review_v3'/(ident+'.json'))
  assert review.get('review'),f'review error {ident}'
  assert review['answer_sha256']==sha(r['reply'].get('answer',''))
  assert review['review']['baseline_answer_sha256']==sha(o['reply'].get('answer',''))
  rows.append({'id':ident,'category':r['category'],'original_question':r['original_question'],'question':r['question'],
   'rewrite_reason':r.get('rewrite_reason'),'expected_output':r.get('expected_output'),
   'before':{k:o.get(k) for k in ('http_status','started_at','completed_at')}|{'answer':o['reply'].get('answer',''),'sources':o['reply'].get('sources',[]),'review':review['review']['before'],'answer_sha256':sha(o['reply'].get('answer',''))},
   'after':{k:r.get(k) for k in ('http_status','started_at','completed_at')}|{'answer':r['reply'].get('answer',''),'sources':r['reply'].get('sources',[]),'review':review['review']['after'],'answer_sha256':sha(r['reply'].get('answer','')),
       'model_label':r['reply'].get('answerability'),'kind':r['reply'].get('answer_kind'),'route':r['reply'].get('evidence_diagnostics',{}).get('route'),
       'crag_audit':r['reply'].get('evidence_diagnostics',{}).get('crag_audit'),'validation':r['reply'].get('validation')},
   'engineering_findings':engineering_findings(o,r),
   'program_checks':{phase:{k:checks[phase][ident].get(k) for k in ('typed_delivery_check','required_measurement_fact_count','missing_returned_citation_ids','missing_mentioned_source_ids','forecast_period_check')} for phase in ('before','after')},
   'blind_order':review['review']['blind_order'],'winner':review['review']['winner'],
   'comparison_reason':display_blind_text(review['review']['comparison_reason'],review['review']['blind_order'])})
 counts={phase:dict(Counter(r[phase]['review']['label'] for r in rows)) for phase in ('before','after')}
 causes={phase:dict(Counter(c for r in rows for c in {u['cause'] for u in r[phase]['review']['unresolved']})) for phase in ('before','after')}
 confirmed={phase:dict(Counter(c for r in rows for c in {u['cause'] for u in r[phase]['review']['unresolved'] if u['confidence']=='confirmed'})) for phase in ('before','after')}
 cats=[]
 for category in dict.fromkeys(r['category'] for r in rows):
  subset=[r for r in rows if r['category']==category]
  cats.append({'category':category,'count':len(subset),**{p:dict(Counter(r[p]['review']['label'] for r in subset)) for p in ('before','after')}})
 both={}
 for phase in ('before','after'):
  records=b['items'] if phase=='before' else a['items'];retrieval_audits=[r['reply'].get('evidence_diagnostics',{}).get('crag_audit') or {} for r in records]
  both[phase]={'local_and_web_executed':sum(x.get('local',{}).get('executed') is True and x.get('web',{}).get('executed') is True for x in retrieval_audits),
    'actually_cited_both':sum((x.get('used_local',0)>0 and x.get('used_web',0)>0) for x in retrieval_audits),
    'corrective_round':sum(bool(x.get('corrective_rounds')) for x in retrieval_audits)}
 capture=load(ROOT/'tmp/crag182_official_sources.json')
 provenance=[{k:r.get(k) for k in ('url','final_url','page_title','published_at','read_status','content_sha256')}|{'characters':len(r.get('content',''))} for r in capture['items']]
 return {'created_at':datetime.now().astimezone().isoformat(),'question_count':160,'expert_accepted':False,'official_source_capture':provenance,
  'program_audits':{phase:{k:v for k,v in audit.items() if k!='items'} for phase,audit in audits.items()},
  'family_delivery_audit':load(ROOT/'tmp/crag182_family_delivery_audit.json'),
  'unseen_variants':[load(p) for p in sorted((ROOT/'tmp/crag182_unseen_family_variants').glob('*.json'))],
  'selected_passage_audits':{phase:{k:v for k,v in audit.items() if k!='items'} for phase,audit in selected.items()},
  'review_method':'同一DeepSeek辅助评阅，按题号哈希盲排A/B；逐字核对评语引用原答。少量引文措辞不一致时，只按至少50字且覆盖80%以上的连续原文锚定，并保留模型原片段，不改变评价结论。不等同独立模型或专家业务验收。',
  'counts':counts,'causes_question_counts_nonexclusive':causes,'confirmed_cause_question_counts_nonexclusive':confirmed,'categories':cats,
  'comparison':dict(Counter(r['winner'] for r in rows)),'retrieval':both,'changes':CHANGES,'remaining_data':DATA_GAPS,
  'engineering_finding_question_counts':dict(Counter(f for r in rows for f in {x['family'] for x in r['engineering_findings']})),
  'before_run':b['run_identity'],'after_run':a['run_identity'],
  'input_hashes':{str(p.relative_to(ROOT)):hashlib.sha256(p.read_bytes()).hexdigest() for p in (before/'answers.json',after/'answers.json',after/'manifest.json')},'items':rows}

def build_html(data,path):
 out=['<!doctype html><html lang="zh-CN"><meta charset="utf-8"><title>CRAG联合检索160题前后实测对照</title><style>body{font:16px/1.75 system-ui,sans-serif;max-width:1250px;margin:30px auto;padding:0 22px;color:#152a39;background:#f7f9fa}h1,h2,h3{line-height:1.4}article,section{background:white;padding:24px;margin:22px 0;border:1px solid #d9e1e5;border-radius:9px}p{white-space:pre-wrap}pre{white-space:pre-wrap;overflow-wrap:anywhere;font:14px/1.7 monospace}summary{cursor:pointer;font-weight:600}table{border-collapse:collapse;width:100%;font-size:14px}td,th{border:1px solid #d2d9dd;text-align:left;vertical-align:top;padding:8px}th{background:#eef2f4}.tablewrap{overflow-x:auto}input,select{padding:9px;font:inherit;max-width:95%}.controls{position:sticky;top:0;background:#f7f9fa;padding:12px;z-index:2}a{color:#006c81}li{overflow-wrap:anywhere}.badge{font-size:14px;color:#42596a}.missing{background:#fff8ec;padding:12px}</style>']
 summary='<table><tr><th>阶段</th><th>完整回答</th><th>部分回答</th><th>未回答</th></tr>'+''.join('<tr><td>'+name+'</td>'+''.join('<td>'+str(data['counts'][phase].get(k,0))+'</td>' for k in ('完整回答','部分回答','未回答'))+'</tr>' for phase,name in [('before','改造前'),('after','改造后')])+'</table>'
 causes='<table><tr><th>根因</th><th>改造前涉及题数</th><th>改造后涉及题数</th></tr>'+''.join('<tr><td>'+html.escape(c)+'</td>'+''.join('<td>'+str(data['causes_question_counts_nonexclusive'][p].get(c,0))+'</td>' for p in ('before','after'))+'</tr>' for c in sorted(set(data['causes_question_counts_nonexclusive']['before'])|set(data['causes_question_counts_nonexclusive']['after'])))+'</table>'
 out+=['<h1>CRAG联合检索160题前后实测对照</h1><p>同一批具体日期和地点，分别调用改造前后系统。以下保留实际返回文字，不以理想答案替换失败答案。完整程度来自证据辅助评阅，另列程序检查，不作为专家正确率。</p>', '<section><h2>结果概览</h2>'+summary+'<p>根因按题统计，可重复归类；“缺数据”与“本轮未检索到”分别保留。确认与推断见逐题原因。</p>'+causes+'</section>']
 out.append('<section><h2>本轮修改</h2><ol>'+''.join('<li>'+html.escape(a+'：'+b)+'<br><small>'+html.escape(c)+'</small></li>' for a,b,c in CHANGES)+'</ol><h2>仍需补齐的业务条件</h2>'+''.join('<p><strong>'+html.escape(a)+'</strong><br>'+html.escape(b+' '+c)+'</p>' for a,b,c in DATA_GAPS)+'</section>')
 out.append('<div class="controls"><input id="search" placeholder="搜索题号、地点、问题或回答"><select id="label"><option value="">全部完整程度</option><option>完整回答</option><option>部分回答</option><option>未回答</option></select><select id="cause"><option value="">全部根因</option>'+''.join('<option>'+html.escape(c)+'</option>' for c in data['causes_question_counts_nonexclusive']['after'])+'</select><span id="count"></span></div>')
 for r in data['items']:
  causes=sorted({u['cause'] for u in r['after']['review']['unresolved']})
  out.append(f'<article data-label="{r["after"]["review"]["label"]}" data-cause="{html.escape("|".join(causes))}"><h2>{r["id"]} {html.escape(r["question"])}</h2><p class="badge">{html.escape(r["category"])}　改造前 {r["before"]["review"]["label"]} → 改造后 {r["after"]["review"]["label"]}</p>')
  out.append('<details><summary>原问题与具体化口径</summary><p>'+html.escape(r['original_question'])+'</p><p>'+html.escape(str(r['rewrite_reason'] or '未改写'))+'</p><p>台账扩展交付要求 '+html.escape(str(r['expected_output']))+'</p></details>')
  out.append('<p>'+html.escape(r['comparison_reason'])+'</p>')
  if r['engineering_findings']:out.append('<div class="missing"><strong>程序与原始记录已确认的原因</strong>'+''.join('<p>'+html.escape(f['cause']+'：'+f['evidence'])+'</p>' for f in r['engineering_findings'])+'</div>')
  if causes:
   out.append('<div class="missing"><strong>仍未答全的内容及根因</strong>'+''.join('<p>'+html.escape(display_blind_text(u['point']+'｜'+u['cause']+'｜'+u['confidence']+'\n依据：'+u['basis']+'\n处理：'+u['next_action'],r['blind_order']))+'</p>' for u in r['after']['review']['unresolved'])+'</div>')
  for phase,title in [('after','新版实际回答'),('before','改造前实际回答')]:
   x=r[phase];out.append('<details'+(' open' if phase=='after' else '')+'><summary>'+title+'</summary>'+answer_html(x['answer'])+sources_html(x['sources'])+'<details><summary>原始返回文字与内容哈希</summary><p>SHA256 '+x['answer_sha256']+'</p><pre class="raw">'+html.escape(x['answer'])+'</pre></details></details>')
  out.append('<details><summary>逐点评阅与检索诊断</summary><pre>'+html.escape(json.dumps({'before':r['before']['review'],'after':r['after']['review'],'retrieval':r['after']['crag_audit'],'validation':r['after']['validation'],'program_checks':r['program_checks']},ensure_ascii=False,indent=2))+'</pre></details></article>')
 out.append('<section><h2>题目之外的变式测试</h2>'+''.join('<h3>'+r['id']+' '+html.escape(r['question'])+'</h3>'+answer_html(r['reply']['answer'])+sources_html(r['reply'].get('sources',[])) for r in data['unseen_variants'])+'</section>')
 out.append('<section><h2>方法与复核边界</h2><p>'+html.escape(data['review_method'])+'</p><p>数值检查只核对程序事实交付、单位、对象、时段及引用；不能证明模式预报误差合格。通用概念题无需强加本地数值，指定业务事实仍需具体适用数据。</p><p>参考方法 <a href="https://arxiv.org/abs/2401.15884">CRAG论文</a> 与 <a href="https://github.com/HuskyInSalt/CRAG">作者代码</a>。本项目独立适配流程，未运行上游训练脚本。</p></section>')
 out.append('<script>const articles=[...document.querySelectorAll("article")];function filter(){let q=document.querySelector("#search").value.toLowerCase(),l=document.querySelector("#label").value,c=document.querySelector("#cause").value,n=0;for(const a of articles){a.hidden=!(a.textContent.toLowerCase().includes(q)&&(!l||a.dataset.label===l)&&(!c||a.dataset.cause.split("|").includes(c)));if(!a.hidden)n++}document.querySelector("#count").textContent=` 显示 ${n}/160 题`}document.querySelectorAll("input,select").forEach(x=>x.addEventListener("input",filter));filter();</script></html>')
 path.write_text('\n'.join(out),encoding='utf-8')
 from html.parser import HTMLParser
 class Raw(HTMLParser):
  def __init__(self):super().__init__(convert_charrefs=True);self.active=False;self.rows=[]
  def handle_starttag(self,t,attrs):
   if t=='pre' and dict(attrs).get('class')=='raw':self.active=True;self.rows.append('')
  def handle_endtag(self,t):
   if t=='pre':self.active=False
  def handle_data(self,s):
   if self.active:self.rows[-1]+=s
 parser=Raw();parser.feed(path.read_text('utf-8'));assert parser.rows==[r[p]['answer'] for r in data['items'] for p in ('after','before')]

def build_docx(data,path):
 from docx import Document
 from docx.shared import Pt
 from build_160_phase16_document import configure,paragraph,table,markdown_answer
 doc=Document();configure(doc);doc.add_paragraph('海洋灾害问答160题实测与改造结果','Title')
 paragraph(doc,'2026年10月19日  改造前 phase1.8.0  改造后 phase1.8.2')
 paragraph(doc,'本报告按同一批具体日期、地点的问题测试新版系统，记录实际回答、未答全的原因及本轮代码修改。逐题附录保留新版原回答；配套HTML保存两轮全部原回答和证据来源。完整程度由DeepSeek辅助盲评及程序核对形成，不等同专家验收或预报准确率。')
 paragraph(doc,'本轮辅助复核结果为完整回答'+str(data['counts']['after'].get('完整回答',0))+'题、部分回答'+str(data['counts']['after'].get('部分回答',0))+'题、未回答'+str(data['counts']['after'].get('未回答',0))+'题。系统仍未达到160题全部完整交付；后续工作需继续补齐适用正文和业务数据，并复核生成遗漏。')
 doc.add_heading('测试结果',level=1)
 table(doc,[['阶段','完整回答','部分回答','未回答']]+[[name]+[str(data['counts'][p].get(k,0)) for k in ('完整回答','部分回答','未回答')] for p,name in [('before','改造前'),('after','改造后')]])
 paragraph(doc,'160题经同一真实HTTP接口逐题调用。测试日期、地点和题目原文一致；不以接口成功、有引用或模型自称完整代替内容检查。两轮回答正文哈希、运行源代码哈希和逐题原始记录均留存。')
 paragraph(doc,'网络返回与资料缓存随时间变化，前后结果反映这两轮实际运行情况，不能把所有差异都归因于代码。曾中断的候选测试另行留存，不混入最终160题统计。')
 table(doc,[['程序检查','改造前','改造后'],['存在程序事实的交付通过题数']+[str(data['program_audits'][p]['typed_delivery_counts'].get('pass',0)) for p in ('before','after')],['所要求的测量事实条数']+[str(data['program_audits'][p]['required_measurement_fact_count']) for p in ('before','after')]])
 paragraph(doc,'程序检查不适用的概念、规则和流程题仍按其必答点与原文核对；通过数值交付检查也不代表SWL或高程基准已获业务核定。')
 table(doc,[['问题类别','题数','改造前完整','改造后完整']]+[[r['category'],str(r['count']),str(r['before'].get('完整回答',0)),str(r['after'].get('完整回答',0))] for r in data['categories']])
 doc.add_heading('未答全原因',level=1)
 paragraph(doc,'下表按涉及的题数统计，同一道题可以同时缺数据和存在检索问题，不能直接相加。confirmed表示辅助评阅给出明确证据，inferred表示推断；逐题附录保留区分。未取得正文不直接认定为全库没有资料。')
 table(doc,[['根因','改造前涉及题数','改造后涉及题数']]+[[c,str(data['causes_question_counts_nonexclusive']['before'].get(c,0)),str(data['causes_question_counts_nonexclusive']['after'].get(c,0))] for c in sorted(set(data['causes_question_counts_nonexclusive']['before'])|set(data['causes_question_counts_nonexclusive']['after']))])
 doc.add_heading('本轮完成的共性修改',level=1)
 table(doc,[['原始记录可直接确认的工程问题','涉及题数']]+[[k,str(v)] for k,v in data['engineering_finding_question_counts'].items()])
 paragraph(doc,'工程归因来自两轮实际路由、变量和校验记录，与辅助评阅的缺项分类分开保留。例如本地NC已有MWD却没有提取，应归为代码或路由问题；MWD方向约定尚未确认，则仍是业务口径缺项。')
 for title,desc,files in CHANGES:paragraph(doc,title,bold=True,keep_next=True);paragraph(doc,desc+' 对应模块：'+files)
 paragraph(doc,'自动回归测试1012项通过，另有102项子测试通过；三个前端脚本测试通过。数值、日期截止、单位拒用、逐时覆盖、双问题路由及历史月度去重均有专门回归案例。')
 doc.add_heading('仍需补齐的数据和业务条件',level=1)
 for title,fields,limit in DATA_GAPS:paragraph(doc,title,bold=True,keep_next=True);paragraph(doc,fields+' '+limit)
 doc.add_heading('逐题实际回答与原因',level=1)
 preservation=[]
 for r in data['items']:
  doc.add_heading(r['id']+' '+r['category'],level=2)
  paragraph(doc,r['question'],bold=True)
  paragraph(doc,'改造前 '+r['before']['review']['label']+' → 改造后 '+r['after']['review']['label']+'。'+r['comparison_reason'])
  for finding in r['engineering_findings']:paragraph(doc,'程序核查：'+finding['cause']+'。'+finding['evidence'])
  if r['program_checks']['after'].get('typed_delivery_check')!='not_applicable':
   paragraph(doc,'程序事实交付检查：'+str(r['program_checks']['after'].get('typed_delivery_check'))+'；需核对的测量事实 '+str(r['program_checks']['after'].get('required_measurement_fact_count',0))+' 项。该检查不代替内容完整程度或预报准确率。',size=10)
  for u in r['after']['review']['unresolved']:paragraph(doc,display_blind_text('未完成：'+u['point']+'；原因：'+u['cause']+'（'+('有明确证据' if u['confidence']=='confirmed' else '待进一步定位')+'）。'+u['basis']+' 处理：'+u['next_action'],r['blind_order']))
  paragraph(doc,'新版实际回答',bold=True,keep_next=True)
  preservation.append({'id':r['id'],**markdown_answer(doc,r['after']['answer'],max_table_cell_characters=450)})
  paragraph(doc,'本题资料来源',bold=True,keep_next=True)
  for n,s in enumerate(r['after']['sources'],1):
   paragraph(doc,f"[资料{s.get('citation_id',n)}] "+str(s.get('title') or s.get('filename') or '未标标题')+'；'+str(s.get('locator') or ''),size=10)
 doc.add_heading('复核方法与原始记录',level=1)
 paragraph(doc,data['review_method'])
 paragraph(doc,'两轮原始数据目录：tmp/qa_full160_crag180_20261019_before 与 tmp/qa_full160_crag182_20261019_after。每题实际回答逐字保存，程序计算依据独立留存。概念解释与定量业务问题分别判断；未请求的本地曲线或实时告警不作为概念题的必答点。')
 paragraph(doc,'参考方法：Corrective Retrieval Augmented Generation，https://arxiv.org/abs/2401.15884；作者代码 https://github.com/HuskyInSalt/CRAG。')
 doc.add_heading('题目之外的变式测试',level=1)
 for r in data['unseen_variants']:
  doc.add_heading(r['id']+' '+r['question'],level=2)
  markdown_answer(doc,r['reply']['answer'],max_table_cell_characters=450)
 doc.add_heading('补充官方来源的正文读取记录',level=1)
 for r in data['official_source_capture']:
  paragraph(doc,str(r.get('page_title') or r['url']),bold=True,keep_next=True)
  paragraph(doc,str(r['url'])+'；正文读取状态 '+str(r['read_status'])+'；提取字符 '+str(r['characters'])+'。',size=10)
 paragraph(doc,'success为成功提取正文，partial为部分正文，failed不能作为已取得正文的依据。2022年预案仅支持注明版本的术语解释，不能冒充现行预案；年度台风风雨灾损与海洋灾害灾损分列，不混用统计口径。')
 from docx.oxml import OxmlElement
 from docx.oxml.ns import qn
 for t in doc.tables:
  for row in t.rows:
   pr=row._tr.get_or_add_trPr()
   if pr.find(qn('w:cantSplit')) is None:pr.append(OxmlElement('w:cantSplit'))
 doc.save(path);save(path.with_suffix('.preservation.json'),{'question_count':160,'items':preservation})

def main():
 p=argparse.ArgumentParser();p.add_argument('--docx',action='store_true');args=p.parse_args()
 before=ROOT/'tmp/qa_full160_crag180_20261019_before';after=ROOT/'tmp/qa_full160_crag182_20261019_after'
 data=collect(before,after);base=ROOT/'exports/CRAG联合检索_160题实际回答与前后对照_20261019'
 save(base.with_suffix('.json'),data);build_html(data,base.with_suffix('.html'))
 if args.docx:build_docx(data,base.with_suffix('.docx'))
 print(json.dumps({'counts':data['counts'],'comparison':data['comparison'],'files':str(base)},ensure_ascii=False))
if __name__=='__main__':main()

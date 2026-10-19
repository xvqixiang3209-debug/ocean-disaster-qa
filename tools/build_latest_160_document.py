"""Compile unchanged live replies and a concrete completion plan for project review."""
import hashlib,json,re
from collections import Counter,OrderedDict
from pathlib import Path
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from build_160_qa_document import inline,field,table,add_answer
from build_160_test_report import link

APP=Path(__file__).resolve().parents[1]
RUN=APP/'tmp/qa_pipeline_after_20261017'
DEST=APP/'exports/海洋灾害160题最新回答与完整正确实现路线_20261017.docx'


def p(doc,text='',size=11,bold=False,keep=False):
    para=doc.add_paragraph();para.paragraph_format.keep_with_next=keep
    inline(para,'**'+text+'**' if bold else text,size)
    return para


def setup():
    doc=Document();s=doc.sections[0]
    s.page_width=Inches(8.5);s.page_height=Inches(11)
    s.top_margin=s.bottom_margin=Inches(.75);s.left_margin=s.right_margin=Inches(1)
    s.header_distance=s.footer_distance=Inches(.3)
    for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','List Bullet','List Number']:
        st=doc.styles[name];st.font.name='Calibri';st.font.color.rgb=RGBColor(0,0,0);st.font.italic=False
        st._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'宋体')
        pp=st._element.get_or_add_pPr()
        for node in list(pp.findall(qn('w:pBdr'))):pp.remove(node)
    normal=doc.styles['Normal'];normal.font.size=Pt(11)
    normal.paragraph_format.line_spacing=1.12;normal.paragraph_format.space_after=Pt(4)
    normal.paragraph_format.widow_control=True
    for name,size,before,after in [('Title',22,0,12),('Subtitle',11,0,10),('Heading 1',16,14,8),('Heading 2',12.5,12,7),('Heading 3',11.5,9,5)]:
        st=doc.styles[name];st.font.size=Pt(size);st.font.bold=name!='Subtitle'
        st.paragraph_format.space_before=Pt(before);st.paragraph_format.space_after=Pt(after)
        st.paragraph_format.keep_with_next=True
    header=s.header.paragraphs[0];inline(header,'海洋灾害160题最新回答与完整正确实现路线',9)
    footer=s.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    inline(footer,'第 ',9);field(footer,'PAGE');inline(footer,' 页  共 ',9);field(footer,'NUMPAGES');inline(footer,' 页',9)
    return doc


def front(doc,report,manifest,rows):
    doc.add_paragraph('海洋灾害160题最新回答与完整正确实现路线',style='Title')
    doc.add_paragraph('2026年10月17日  phase1.4.1  最新实测汇编',style='Subtitle')
    p(doc,'本文收录回答链路修改后的160道具体问题、系统完整原答、引用来源和系统记录的缺项，并说明把当前结果提升为完整、正确业务答案的开发路线。正文没有另写一套理想答案替换系统输出，便于逐题定位和复核。')
    p(doc,'本次160题均取得DeepSeek回复，原52题生成失败已消除。但这一结果只证明回答链路运行成功。模型自报6题完整覆盖、154题部分覆盖；其中16题没有可引用来源，只给出通用解释或处理路径。150题使用了通用解释补充。完整性和正确性尚未完成专家逐条验收。')
    p(doc,'我们接下来的目标是逐题完成约定的必答内容，并使事实、数字、规则、时间和地点与可复核依据一致。没有正式数据的当前结论、尚未发布的命令以及无法确定的未来结果，不能通过补写文字获得正确答案；正确回答要说明已知结果、不确定性和可查询的下一步。')
    doc.add_paragraph('如何使用这份文档',style='Heading 1')
    p(doc,'先阅读实现路线，再按编号或类别查题。每题先列实测问题和系统原答，随后列来源、系统自报缺项和原需求目标。资料编号仅在本题内对应；系统原答中的日期、台风和数值均属于当时场景，不作为阅读此文时的实时状态。')
    p(doc,'本次采集时间为'+report['started_at']+'至'+report['completed_at']+'，模型为'+report['model']+'。160题沿用前次实测的同一清单，每题独立提问，不带上一题历史。连续追问另做真实接口验证，已正确承接“风暴增水”与“它和天文大潮叠加会怎样”。',10)
    p(doc,'“这次台风”主要具体化为'+manifest['storm']['name']+'（'+manifest['storm']['sid']+'）；地点和未来预报区间沿用原清单。原题与实测题不同时，两者均列出。“假设”“情景”不表示事件已经发生。问答原始记录的场景摘要为 '+report['manifest_sha256']+'。',9.5)
    doc.add_paragraph('当前结果应该怎样理解',style='Heading 1')
    table(doc,[['检查项目','修改前','修改后'],['DeepSeek返回回答','87题','160题'],['生成或引用校验失败','52题','0题'],['仅返回依据不足模式','21题','0题'],['模型自报完整覆盖','3题','6题'],['模型自报部分覆盖','84题','154题'],['当前无可引用来源','本项未按同口径复核','16题'],['专家逐题验收','未完成','未完成']],[3.2,1.65,1.65])
    p(doc,'“仅返回依据不足模式为0”表示系统改为提供通用解释，不表示缺失的业务数据已经补齐。当前6题自报完整也不代表6题已经正确验收；例如台账若要求月度频次表，只有季节性说明仍需补齐统计结果。',10)
    p(doc,'154题部分覆盖的原因包含真实缺数、检索或路由不足、校验退回以及答题范围扩张。Q-001询问灾种区别和居民关注重点，却被加入“当前实况”缺项；基础概念题不应因此无限延长到实时查询。R-006已有业务查询ready状态，却无最终引用，需检查查询结果能否真正回答“验潮站”，以及生成和校验过程。缺项标签本身也要复核。')
    doc.add_paragraph('完整正确答案的实现路线',style='Heading 1')
    plan=[
      ('先逐题确定必答内容','我们将160题拆为可检查的子问题，记录需要的答案字段、地点、时间、业务对象、来源和完成标准。定义题只要求定义、区别和适当例子；询问当前状态时才加入实时数据。原台账要求的月度表、站点排名、地图或导出文件单独列为交付项，避免仅有说明文字就记为完成。','完成标志是160题都有稳定的必答清单；原题、具体场景和验收范围一致，不增加用户没有问的实时问题。'),
      ('先修已有资料和数据的利用','我们逐题检查已导入材料、检索结果、业务调用和引用链。把权威定义、常用术语、避险知识整理为带来源和版本的基础知识条目；按每个子问题分别检索、重排和取正文。方法题不再用只有同名词的预报简报替代技术资料。对已有数据但被校验退回的题，检查实际结论和证据对应关系后修复。','完成标志是基础知识题能直接完整回答，出处支持具体结论；静态题不会因缺少当前实况被判定未完成。'),
      ('把业务查询做成可组合的工具','我们保留现有预报查询，补充按站点、岸段、空间位置、时段、变量和批次调用的统一工具，并向问题规划器提供明确的输入输出。检索工具解决文献问题，数据库和计算程序解决数字问题。模型选择工具、解释结果，不自行推算缺少口径的数据。每个子问题记录选用的工具和结果。','完成标志是查询可复现，数值包含对象、单位、基准、起报和有效时间；缺数不当零，岸段不冒充验潮站，阵风不冒充平均风。'),
      ('接入真正缺少的业务数据','我们按题号建立接入清单，优先补齐正式预警与公告、实时观测、站点预报、历史观测和灾损；再补海堤、人口、淹没范围、安置容量、交通和抢修等台账。NC解析结果进入结构化业务库并保留原文件、变量元数据和批次，知识库保存说明、标准和引用入口。','完成标志是每个业务子问题有真实的数据提供方、更新方式、字段和缺测状态；没有接口的公告来源能够检索并核对正文。'),
      ('核定影响数值结论的口径','我们把SWL含义、天文潮高程基准、单位、时区、站点关联、岸段阈值和无效值整理为一张业务口径表，以数据说明或提供方确认替换当前试算假设。通过已知站点和时刻做人工对照，确认叠加和阈值判断结果。','完成标志是关键口径有出处、版本与确认记录，典型时刻的程序结果与人工对照一致；未核定项继续明确标为试算，不升为正式结论。'),
      ('按子问题检查是否答全','每个子问题先取得相应资料或工具结果。首次证据不足时，对缺少的具体部分补充检索和官方正文读取；达到设定次数后给出有依据的部分并准确说明剩余未知。生成时先给结论，再按问题组织解释、数据或步骤，删去重复通用段落和无关历史数值。','完成标志是每个必答点都对应答案位置和证据；没有因已回答一个子问题就遗漏其他部分，也不以重复免责声明替代回答。'),
      ('分别验收完整性与正确性','我们建立160题的参考答案或期望结果，按必答点覆盖、事实一致、数字与单位、规则版本和地域、来源对应、时间有效性逐项检查。动态题固定输入快照与查询时刻，另测最新数据更新。原题之外再测改写问题、跨轮追问和缺数据边界，避免只记住160个固定说法。模型评分用于辅助筛查，最终业务判断由专业人员对照证据确认。','完成标志是每题保留参考结果、实际回答、证据和验收记录；未通过项有具体原因，修复后复测，生成成功不再直接记为正确。'),
    ]
    for i,(title,action,done) in enumerate(plan,1):
        doc.add_paragraph(str(i)+' '+title,style='Heading 2');p(doc,action);p(doc,done,10)
    doc.add_paragraph('优先完成的工作与对应题目',style='Heading 1')
    matrix=[['工作','典型题号','完成后可得到的结果'],
      ['答题范围与基础知识','Q-001 Q-005 Q-039 Q-048 R-073','概念和一般方法完整回答，引用对应，不额外要求实时状态。'],
      ['站点预报与工具结果利用','Q-012 R-005 R-006 R-033','按站点或岸段、空间位置和时段返回正确数值，分清业务对象及风种。'],
      ['正式预警与行政公告','Q-015 R-019 R-020','按地区、发布时间、有效时段和发文主体解释预警、五停和响应决定。'],
      ['淹没风险与现场台账','Q-023 R-015','结合地形、海堤、边界水位及现场台账回答具体风险，不靠浪高或增水直接判断住户进水。'],
      ['统计与预报检验','Q-003 Q-044 R-045','完整统计表、比较指标和观测配对检验结果可复算。'],
      ['现行政策与标准版本','R-047 R-059','明确适用地区、发布实施及替代关系，归档版本不直接当作现行版本。']]
    table(doc,matrix,[1.55,1.6,3.35])
    p(doc,'上述题号用于定位典型问题，各工作可以重叠。具体“是否确实缺数据”以逐题资料核查为准，不按模型自报缺项直接采购或重复导入。')
    doc.add_paragraph('重新核对发现的典型问题',style='Heading 1')
    for title,text in [
      ('Q003 自报完整但没有月度频次表','本题回答台风活动季节和重点月份，但原需求目标包含月度频次表及统计口径。我们需要从完整路径资料按明确的活动、影响或登陆定义统计，列出各月数量及样本年份；季节说明不能代替统计表。'),
      ('Q012 自报完整但要素尚不齐全','本题给出SSW假设阵风和WAV时段极值。原需求还包含风向、测高与平均时段、有效波高与周期波向、起报和有效时间等信息。未接入或元数据不支持的变量要列出，已支持的变量则补查；不能只凭数值存在就判断全部完成。'),
      ('R033 最大值的覆盖范围需要限定','程序本次请求198个岸段，可计算173个，25个缺数。原答首句使用“浙江岸段最大”的范围过宽；能够证明的是173个可计算岸段中的最大值，不能排除缺数岸段更大。完整答案还需要模式版本、分辨率和覆盖范围说明。'),
      ('R036 已有增水计算没有正确调用','本题的天文潮和增水叠加问题只进入天文潮分支，随后原答称缺增水过程；现有coastal计算已支持同刻叠加。我们先修多子任务路由，取得一致位置和时段的过程，再核验叠加和最不利时刻，不把该问题直接列为用户未提供增水文件。'),
      ('Q051 系统能力说明没有被命中','系统已有运行配置说明，但本题走资料检索并称缺少系统说明。我们扩展意图识别，直接读取实际配置和已接入工具，重新核对能做与不能做的能力描述。R-028的核定工作安排被误当潮位数值查询、Q-014相似台风没有使用现有相似检索，也优先修复。'),
    ]:
        doc.add_paragraph(title,style='Heading 2');p(doc,text)
    p(doc,'以上是实测输出和代码路径核对中发现的问题，不是对其他题已正确的保证。下文仍原样保留这些题的系统回答，便于对照修改；没有将编辑者指出的修正悄悄替换进原答。')
    doc.add_paragraph('需要业务方集中确认的内容',style='Heading 1')
    p(doc,'我们可以先完成检索、工具调用、覆盖检查、文档和测试代码。业务方集中确认三类内容即可：关键数值口径与数据说明；正式预警、实时观测及现场台账的提供渠道；160题的必答范围和专业参考结果。确认发生在资料和规则维护阶段，普通用户提问时无需逐条审核。')
    p(doc,'对尚未发布的复航时间、行政决定或未来登陆结果，完整的回答可以是“当前正式信息说明了什么、哪些尚未确定、后续到哪里核查”。系统不会为了填满答案而编造确定结论，也不承诺零幻觉；可靠性由可追溯数据、计算校验和持续逐题验收共同提高。')
    doc.add_paragraph('试算与阅读限制',style='Heading 1')
    p(doc,'现有数值沿用用户授权的假设：SWL按纯增水并以米处理，天文潮按已订正至1985基准处理，岸段阈值按厘米及1985基准解释，未声明时区按北京时间解释。无效阈值不补造。假设不等于已核定，数值合理性不能独自证明基准正确。')
    p(doc,'本次239项自动测试和21项子测试通过，证明所测程序行为；引用存在和数字匹配也不能独自证明语义正确。本文不把历史公报、旧预警、模型通用知识或试算数值提升为当前正式决定。系统记录的缺项在各题后原样列出，可能包含超出题意的附加要求，仍需按必答清单复核。')
    doc.add_paragraph('技术依据',style='Heading 1')
    p(doc,'DeepSeek官方工具调用文档说明，模型产生工具调用请求，具体函数由应用执行；因此接入大模型不会自动获得业务数据。完整性、正确性与依据一致性分别检查的做法，也见Microsoft的RAG评测说明。以下仅作为技术路线参考，不作为海洋业务依据。',10)
    a=p(doc,'DeepSeek工具调用说明  ',10);link(a,'https://api-docs.deepseek.com/guides/tool_calls/','官方文档')
    a=p(doc,'Microsoft RAG评测说明  ',10);link(a,'https://learn.microsoft.com/en-us/azure/foundry/concepts/evaluation-evaluators/rag-evaluators','官方文档')
    doc.add_paragraph('问题类别索引',style='Heading 1')
    groups=OrderedDict()
    for row in rows:groups.setdefault(row['category'],[]).append(row)
    table(doc,[['类别','题数']]+[[cat,str(len(items))] for cat,items in groups.items()],[5.8,.7])
    doc.add_paragraph('没有可引用来源的题目',style='Heading 2')
    ids=[r['id'] for r in rows if not r['result'].get('rag_sources') and not r['result'].get('web_sources')]
    p(doc,'、'.join(ids),10)
    p(doc,'这些题仍有通用解释，但不能据此记为已完成事实验证；其中也可能有程序已取得数据但未形成有效引用的情况。')
    return groups


def main():
    report=json.loads((RUN/'answers.json').read_text(encoding='utf-8'))
    manifest=json.loads((RUN/'manifest.json').read_text(encoding='utf-8'));rows=report['items']
    assert len(rows)==160 and len({x['id'] for x in rows})==160
    doc=setup();groups=front(doc,report,manifest,rows)
    for i,(cat,items) in enumerate(groups.items(),1):
        doc.add_paragraph(f'{i:02d} {cat}',style='Heading 1')
        for item in items:
            result=item['result'];meta=result.get('answer_meta',{})
            doc.add_paragraph(item['id']+' '+re.sub(r'[^\w\u4e00-\u9fff ]',' ',item.get('subrequirement','')).strip(),style='Heading 2')
            p(doc,'实测问题  '+item['question'],11,True,True)
            if item['question']!=item['original_question']:p(doc,'原题  '+item['original_question'],9.5,keep=True)
            sources=sorted(result.get('rag_sources',[])+result.get('web_sources',[]),key=lambda s:s.get('citation_id',0))
            label='完整覆盖' if meta.get('answerability')=='full' else '部分覆盖'
            if not sources:label+='  无可引用来源'
            if meta.get('business_trial'):label+='  假设口径试算'
            p(doc,'模型自报状态  '+label+'  采集时间  '+item['collected_at'],9.5,keep=True)
            p(doc,'系统完整原答',10,True,True)
            before=len(doc.paragraphs)
            add_answer(doc,result['answer'])
            # Bold-only and short numbered answer headings belong with the
            # paragraph they introduce, including on a page boundary.
            for paragraph in doc.paragraphs[before:]:
                text=paragraph.text.strip()
                bold_only=bool(text) and all(run.bold for run in paragraph.runs if run.text.strip())
                short_numbered=len(text)<=45 and re.match(r'^(?:\d+[.．、]|[一二三四五六七八九十]+[、．])\s*',text)
                if (bold_only and len(text)<=80) or short_numbered:
                    paragraph.paragraph_format.keep_with_next=True
            if sources:
                p(doc,'本题引用来源',10,True,True)
                for src in sources:
                    name=src.get('title') or src.get('filename') or '未命名来源'
                    text='['+str(src.get('citation') or '资料'+str(src.get('citation_id','')))+'] '+name
                    if src.get('locator'):text+='  '+str(src['locator'])
                    if src.get('review_label'):text+='  '+src['review_label']
                    para=p(doc,text,9.5);para.paragraph_format.space_after=Pt(2)
                    if src.get('url'):inline(para,'  ',9.5);link(para,src['url'])
            else:p(doc,'本题没有返回可引用来源。上文通用解释未完成逐项事实核验。',9.5)
            gaps=list(dict.fromkeys(meta.get('missing_information') or []))
            if gaps:
                p(doc,'系统记录的待核实项',10,True,True)
                for gap in gaps:p(doc,'• '+gap,9.5)
            p(doc,'原需求目标  '+item.get('visible_result','未注明'),9.5)
            if meta.get('business_evidence_id'):p(doc,'本次业务查询记录  '+meta['business_evidence_id'],9)
    update=OxmlElement('w:updateFields');update.set(qn('w:val'),'true');doc.settings._element.append(update)
    doc.core_properties.title='海洋灾害160题最新回答与完整正确实现路线'
    doc.core_properties.subject='回答链路修改后实际输出与逐题完善方法'
    doc.core_properties.author='海洋灾害问答项目'
    DEST.parent.mkdir(parents=True,exist_ok=True);doc.save(DEST)
    print(json.dumps({'output':str(DEST),'questions':160,'source_run':str(RUN),'docx_sha256':hashlib.sha256(DEST.read_bytes()).hexdigest()},ensure_ascii=False))


if __name__=='__main__':main()

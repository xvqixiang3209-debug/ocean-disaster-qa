"""Recover original numeric cells into a separate, queryable review archive.

This archive never changes knowledge-source review or transmission permission.
Raw numbers are not silently given NC units, a vertical datum, or a missing date.
"""
from __future__ import annotations

from contextlib import closing
from datetime import datetime, date, timezone, timedelta
from pathlib import Path
import hashlib
import json
import math
import re
import sqlite3

PARSER_VERSION = 'numeric-review-v1'
INVALID = {-999, -9999, -99999, 9999, 99999}
COLOURS = ('blue', 'yellow', 'orange', 'red')
LEVELS = {'blue':'BlueLevel', 'yellow':'YellowLevel', 'orange':'OrangeLevel', 'red':'RedLevel'}


def digest(path):
    sha = hashlib.sha256()
    with Path(path).open('rb') as handle:
        for block in iter(lambda: handle.read(1024*1024), b''):
            sha.update(block)
    return sha.hexdigest()


def dump(value):
    return json.dumps(value, ensure_ascii=False, allow_nan=False, separators=(',', ':'))


def name_key(value):
    return re.sub(r'\s+', '', str(value or '')).casefold()


def coast_key(value):
    # Parenthetical station annotations and * flags are not the coastal name.
    # County names are retained: no fuzzy transfer between jurisdictions.
    return name_key(re.split(r'[（(]',str(value or ''))[0].rstrip('*'))


def _coast_candidates(rows):
    # A full ADMC coastal label and SQL coastal name are specific identities.
    # A generic TOWN can describe several segments and must not override them.
    if not rows:return []
    weight=lambda row:1 if row['source_version']=='geojson_town' else 2
    best=max(weight(row) for row in rows)
    return sorted({row['coast_id'] for row in rows if weight(row)==best})


def number(value):
    """Preserve zero and signed values; keep missing/formulas distinguishable."""
    if value is None or isinstance(value, str) and not value.strip():
        return None, 'missing'
    if isinstance(value, bool):
        return None, 'non_numeric'
    try:
        result = float(value)
    except (ValueError, TypeError):
        return None, 'formula_unresolved' if str(value).startswith('=') else 'non_numeric'
    if not math.isfinite(result):
        return None, 'nonfinite'
    return (None, 'invalid_sentinel') if result in INVALID else (result, 'valid')


def _json_value(value):
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, float) and not math.isfinite(value):
        return str(value)
    return value if isinstance(value, (str, int, float, bool)) or value is None else str(value)


def time_value(value, profile=None):
    """Only full dated values yield dates. DD/HH:mm remains a searchable token."""
    profile = profile or {}
    raw = _json_value(value)
    text = str(raw or '').strip()
    parsed = value if isinstance(value, datetime) else None
    if parsed is None and re.fullmatch(r'\d{4}[-/]\d{1,2}[-/]\d{1,2}(?:[ T].*)?', text):
        try:
            parsed = datetime.fromisoformat(text.replace('/', '-').replace('Z', '+00:00'))
        except ValueError:
            pass
    # Explicit per-token dates are assumptions, never inferred from file names.
    explicit = (profile.get('time_map') or {}).get(text)
    if explicit:
        try:
            assumed = datetime.fromisoformat(str(explicit).replace('Z', '+00:00'))
            if assumed.tzinfo is None:
                return text, None, 'assumption_missing_timezone'
            return text, assumed.astimezone(timezone.utc).isoformat(), 'assumed'
        except ValueError:
            return text, None, 'invalid_assumed_time'
    if parsed is not None:
        if parsed.tzinfo is not None:
            return text, parsed.astimezone(timezone.utc).isoformat(), 'explicit'
        offset = profile.get('timezone_offset_hours')
        if isinstance(offset, (int, float)) and -14 <= offset <= 14:
            return text, parsed.replace(tzinfo=timezone(timedelta(hours=offset))).astimezone(timezone.utc).isoformat(), 'assumed'
        return text, None, 'date_known_timezone_unknown'
    if re.fullmatch(r'\d{1,2}/\d{1,2}:\d{2}(?::\d{2})?', text):
        return text, None, 'day_clock_only'
    return text, None, 'missing' if not text else 'unresolved'


def geometry_metadata(geometry):
    pairs=[];errors=[]
    def visit(item):
        if isinstance(item,(list,tuple)) and len(item)>=2 and not isinstance(item[0],(list,tuple)):
            try:
                lon,lat=float(item[0]),float(item[1])
                if not math.isfinite(lon+lat) or not -180<=lon<=180 or not -90<=lat<=90:
                    raise ValueError()
                pairs.append((lon,lat))
            except (TypeError,ValueError):errors.append('invalid_coordinate_pair')
        elif isinstance(item,(list,tuple)):
            for part in item:visit(part)
    visit((geometry or {}).get('coordinates',[]))
    return {'geometry_type':(geometry or {}).get('type'),'coordinate_pair_count':len(pairs),
            'coordinate_issues':list(set(errors)),
            'bbox':[min(p[0] for p in pairs),min(p[1] for p in pairs),max(p[0] for p in pairs),max(p[1] for p in pairs)] if pairs else None}


class NumericArchive:
    def __init__(self, database, *, read_only=False):
        self.database = Path(database).resolve()
        self.read_only = bool(read_only)
        if self.database.name in {'knowledge.sqlite3', 'typhoon_rag.sqlite3', 'typhoon_history.sqlite3', 'forecast.sqlite3'}:
            raise ValueError('数值复核归档必须使用独立数据库，不能写业务库')
        self.aliases = {}
        self.coast_aliases = None
        self.registry_digest = None
        if self.read_only:
            if not self.database.is_file():raise FileNotFoundError('数值归档尚未建立')
            return
        self.database.parent.mkdir(parents=True, exist_ok=True)
        with self.connect() as db:
            db.executescript('''
            CREATE TABLE IF NOT EXISTS numeric_files(
                file_id TEXT PRIMARY KEY, source_id TEXT,filename TEXT,sha256 TEXT,bytes INTEGER,
                parser_version TEXT,profile TEXT,kind TEXT,state TEXT,error_type TEXT,
                imported_at TEXT,filename_issue_token TEXT);
            CREATE TABLE IF NOT EXISTS numeric_sheets(
                sheet_id INTEGER PRIMARY KEY,file_id TEXT REFERENCES numeric_files(file_id) ON DELETE CASCADE,
                name TEXT,header_row INTEGER,time_col INTEGER,row_count INTEGER,column_count INTEGER,
                unit TEXT,unit_status TEXT,vertical_datum TEXT,datum_status TEXT,ensemble_statistic TEXT);
            CREATE TABLE IF NOT EXISTS numeric_columns(
                sheet_id INTEGER REFERENCES numeric_sheets(sheet_id) ON DELETE CASCADE,
                col_index INTEGER,label TEXT,station_id TEXT,mapping_status TEXT,candidates TEXT,
                PRIMARY KEY(sheet_id,col_index));
            CREATE TABLE IF NOT EXISTS numeric_rows(
                sheet_id INTEGER REFERENCES numeric_sheets(sheet_id) ON DELETE CASCADE,row_index INTEGER,
                time_raw TEXT,valid_time TEXT,time_status TEXT,PRIMARY KEY(sheet_id,row_index));
            CREATE TABLE IF NOT EXISTS numeric_cells(
                sheet_id INTEGER REFERENCES numeric_sheets(sheet_id) ON DELETE CASCADE,
                row_index INTEGER,col_index INTEGER,value REAL,raw TEXT,quality TEXT,formula TEXT,
                PRIMARY KEY(sheet_id,row_index,col_index)) WITHOUT ROWID;
            CREATE INDEX IF NOT EXISTS numeric_station ON numeric_columns(label,station_id);
            CREATE INDEX IF NOT EXISTS numeric_filename ON numeric_files(filename);
            CREATE INDEX IF NOT EXISTS numeric_file_sheets ON numeric_sheets(file_id,sheet_id);
            CREATE INDEX IF NOT EXISTS numeric_cell_station ON numeric_cells(sheet_id,col_index,row_index);
            CREATE INDEX IF NOT EXISTS numeric_time ON numeric_rows(valid_time,time_raw,sheet_id);
            CREATE TABLE IF NOT EXISTS numeric_reference_sources(
                role TEXT PRIMARY KEY,filename TEXT,sha256 TEXT,metadata TEXT);
            CREATE TABLE IF NOT EXISTS numeric_stations(
                station_id TEXT PRIMARY KEY,name TEXT,data_file TEXT,base85 REAL,metadata TEXT);
            CREATE TABLE IF NOT EXISTS numeric_station_aliases(
                alias TEXT,station_id TEXT,PRIMARY KEY(alias,station_id));
            CREATE TABLE IF NOT EXISTS numeric_coasts(
                version TEXT,coast_id TEXT,name TEXT,reference_station TEXT,geometry TEXT,metadata TEXT,
                PRIMARY KEY(version,coast_id));
            CREATE TABLE IF NOT EXISTS numeric_thresholds(
                version TEXT,coast_id TEXT,colour TEXT,raw TEXT,value REAL,quality TEXT,unit TEXT,
                unit_status TEXT,PRIMARY KEY(version,coast_id,colour));
            CREATE TABLE IF NOT EXISTS numeric_coast_conflicts(
                coast_id TEXT,field TEXT,geojson_value TEXT,sql_value TEXT,
                PRIMARY KEY(coast_id,field));
            ''')
            columns={row[1] for row in db.execute('PRAGMA table_info(numeric_columns)')}
            for col,definition in [('coast_id','TEXT'),('coast_mapping_status',"TEXT DEFAULT 'unmapped'"),('coast_candidates',"TEXT DEFAULT '[]'")]:
                if col not in columns:db.execute('ALTER TABLE numeric_columns ADD COLUMN '+col+' '+definition)
            db.execute('CREATE INDEX IF NOT EXISTS numeric_coast_columns ON numeric_columns(coast_id,sheet_id,col_index)')
            db.execute('''CREATE TABLE IF NOT EXISTS numeric_coast_aliases(
                alias TEXT,coast_id TEXT,source_version TEXT,PRIMARY KEY(alias,coast_id,source_version))''')

    def connect(self):
        db = (sqlite3.connect(self.database.as_uri()+'?mode=ro',uri=True,timeout=60)
              if self.read_only else sqlite3.connect(self.database, timeout=60))
        db.row_factory = sqlite3.Row
        db.execute('PRAGMA foreign_keys=ON')
        if self.read_only:db.execute('PRAGMA query_only=ON')
        return db

    def _writable(self):
        if self.read_only:raise ValueError('数值归档以只读模式打开，不能恢复或变更来源')

    def load_station_registry(self, site_sql, names_txt=None):
        self._writable()
        from forecast_runtime.storage import read_sql
        rows = read_sql(site_sql, 'm_fcst_site')
        aliases = {}
        by_code = {}
        def add(label, station_id):
            if label:
                aliases.setdefault(name_key(label), set()).add(str(station_id))
        with self.connect() as db:
            db.execute('DELETE FROM numeric_stations')
            for row in rows:
                sid = str(row['StationID'])
                add(row['StationName'], sid)
                add(row.get('DataFile'), sid)
                if row.get('DataFile'):
                    by_code.setdefault(name_key(row['DataFile']), set()).add(sid)
                db.execute('INSERT INTO numeric_stations VALUES(?,?,?,?,?)',
                    (sid,row['StationName'],row.get('DataFile'),row.get('Base85'),
                     dump({'is_enable':row.get('IsEnable'),'base85_status':'source_value_not_datum_confirmation'})))
            db.execute('INSERT OR REPLACE INTO numeric_reference_sources VALUES(?,?,?,?)',
                ('stations',Path(site_sql).name,digest(site_sql),dump({'row_count':len(rows)})))
            if names_txt:
                text = Path(names_txt).read_text(encoding='utf-8-sig')
                for line in text.splitlines():
                    bits = line.split(None, 1)
                    if len(bits) != 2:
                        continue
                    codes = by_code.get(name_key(bits[0]), set())
                    label = re.sub(r'（.*?）|\(.*?\)', '', bits[1]).strip()
                    label = re.sub(r'(?:国家站|基础站|简易站|海洋站|水利)$', '', label)
                    for part in label.split('/'):
                        for sid in codes:
                            add(part, sid)
                db.execute('INSERT OR REPLACE INTO numeric_reference_sources VALUES(?,?,?,?)',
                    ('station_names',Path(names_txt).name,digest(names_txt),dump({'method':'explicit_code_name_pairs'})))
            db.execute('DELETE FROM numeric_station_aliases')
            db.executemany('INSERT INTO numeric_station_aliases VALUES(?,?)',
                [(label,sid) for label,ids in aliases.items() for sid in ids])
        self.aliases = {key: sorted(value) for key,value in aliases.items()}
        self.registry_digest = digest(site_sql)+(':'+digest(names_txt) if names_txt else '')
        return {'station_rows':len(rows),'alias_count':len(self.aliases)}

    def map_station(self, label):
        if not self.aliases:
            with self.connect() as db:
                rows=db.execute('SELECT alias,station_id FROM numeric_station_aliases').fetchall()
            for row in rows:self.aliases.setdefault(row['alias'],[]).append(row['station_id'])
        candidates = self.aliases.get(name_key(label), [])
        if not candidates and str(label).endswith('站'):
            candidates = self.aliases.get(name_key(str(label)[:-1]), [])
        return (candidates[0] if len(candidates)==1 else None,
                'unique_source_name' if len(candidates)==1 else 'ambiguous' if candidates else 'unmapped', candidates)

    def recover_workbook(self, path, *, source_id=None, filename=None, expected_sha256=None, profile=None):
        self._writable()
        from openpyxl import load_workbook
        from openpyxl.utils import get_column_letter
        path = Path(path)
        sha = digest(path)
        if expected_sha256 and sha != expected_sha256:
            raise ValueError('源文件SHA256与来源登记不一致')
        filename = filename or path.name
        file_id = source_id or sha
        profile = {**(profile or {}),'station_registry_digest':self.registry_digest}
        with self.connect() as db:
            existing = db.execute('SELECT * FROM numeric_files WHERE file_id=?', (file_id,)).fetchone()
            if existing and existing['sha256']==sha and existing['profile']==dump(profile) and existing['parser_version']==PARSER_VERSION and existing['state']=='ready':
                return {'file_id':file_id,'filename':filename,'state':'existing','sha256':sha}
        wb = load_workbook(path, read_only=True, data_only=False, keep_links=False)
        cached = None
        counts = {'cells':0,'valid':0,'missing':0,'invalid_sentinel':0,'formula_unresolved':0,'non_numeric':0,'nonfinite':0,'formula_cached':0}
        try:
            with self.connect() as db:
                db.execute('DELETE FROM numeric_files WHERE file_id=?',(file_id,))
                issue = re.search(r'_(20\d{8,10})\.xlsx$', filename, re.I)
                db.execute('INSERT INTO numeric_files VALUES(?,?,?,?,?,?,?,?,?,?,?,?)',
                    (file_id,source_id,filename,sha,path.stat().st_size,PARSER_VERSION,dump(profile),'workbook','ready',None,
                     datetime.now(timezone.utc).isoformat(),issue[1] if issue else None))
                for sheet in wb:
                    iterator = sheet.iter_rows()
                    first = next(iterator, ())
                    if not first:
                        continue
                    headers = [str(cell.value).strip() if cell.value is not None else '' for cell in first]
                    time_col = next((index for index,label in enumerate(headers) if label in {'时间','日期','time','Time'}), None)
                    if time_col is None:
                        raise ValueError('工作表未识别到显式时间列表头')
                    unit = str(profile.get('unit') or 'unknown')
                    datum = str(profile.get('vertical_datum') or 'unknown')
                    cur=db.execute('INSERT INTO numeric_sheets(file_id,name,header_row,time_col,row_count,column_count,unit,unit_status,vertical_datum,datum_status,ensemble_statistic) VALUES(?,?,?,?,?,?,?,?,?,?,?)',
                        (file_id,sheet.title,1,time_col+1,0,len(headers),unit,'assumed' if unit!='unknown' else 'unknown',datum,
                         'assumed' if datum!='unknown' else 'unknown',str(profile.get('ensemble_statistic') or 'unknown')))
                    sheet_id = cur.lastrowid
                    columns = [i for i,label in enumerate(headers) if label and i!=time_col]
                    for col in columns:
                        sid,state,candidates = self.map_station(headers[col])
                        coastal=self.map_coast(headers[col])
                        coastal_state='unique_source_coast_name' if len(coastal)==1 else 'ambiguous' if coastal else 'unmapped'
                        db.execute('INSERT INTO numeric_columns(sheet_id,col_index,label,station_id,mapping_status,candidates,coast_id,coast_mapping_status,coast_candidates) VALUES(?,?,?,?,?,?,?,?,?)',
                            (sheet_id,col+1,headers[col],sid,state,dump(candidates),coastal[0] if len(coastal)==1 else None,coastal_state,dump(coastal)))
                    row_count = 0
                    batch=[]
                    for row_index,row in enumerate(iterator,2):
                        if not any(cell.value is not None for cell in row):
                            continue
                        token,valid,ts = time_value(row[time_col].value if time_col<len(row) else None,profile)
                        db.execute('INSERT INTO numeric_rows VALUES(?,?,?,?,?)',(sheet_id,row_index,token,valid,ts))
                        row_count+=1
                        for col in columns:
                            cell = row[col] if col<len(row) else None
                            raw = cell.value if cell else None
                            formula = str(raw) if cell and cell.data_type=='f' else None
                            value,quality = number(raw)
                            if formula:
                                if cached is None:
                                    cached=load_workbook(path,read_only=True,data_only=True,keep_links=False)
                                cache=cached[sheet.title].cell(row_index,col+1).value
                                value,quality = number(cache)
                                quality='formula_cached' if quality=='valid' else 'formula_unresolved'
                            counts['cells']+=1;counts[quality]=counts.get(quality,0)+1
                            batch.append((sheet_id,row_index,col+1,value,dump(_json_value(raw)),quality,formula))
                        if len(batch)>=10000:
                            db.executemany('INSERT INTO numeric_cells VALUES(?,?,?,?,?,?,?)',batch);batch=[]
                    if batch:db.executemany('INSERT INTO numeric_cells VALUES(?,?,?,?,?,?,?)',batch)
                    db.execute('UPDATE numeric_sheets SET row_count=? WHERE sheet_id=?',(row_count,sheet_id))
            return {'file_id':file_id,'filename':filename,'state':'recovered','sha256':sha,**counts}
        finally:
            wb.close()
            if cached:cached.close()

    def query(self, *, file=None, station=None, coast=None, time_token=None, start=None, end=None, allow_assumed=False, limit=100):
        from openpyxl.utils import get_column_letter
        limit=min(10000,max(1,int(limit)))
        where=[];args=[]
        if file:
            # Resolve tiny file metadata first; avoid a full cell-table scan
            # caused by an OR across filenames and IDs in the ten-million grid.
            with self.connect() as db:
                file_ids=[r[0] for r in db.execute('SELECT file_id FROM numeric_files WHERE filename=? OR file_id=? OR source_id=?',(file,)*3)]
            where.append('f.file_id IN ('+(','.join('?' for _ in file_ids) if file_ids else 'NULL')+')');args.extend(file_ids)
        if station:
            _,_,candidates=self.map_station(station)
            alternatives=list(dict.fromkeys([str(station),*candidates]))
            coast_ids=self.map_coast(station)
            where.append('(c.label=? OR c.station_id IN ('+','.join('?' for _ in alternatives)+')'+
                         (' OR c.coast_id IN ('+','.join('?' for _ in coast_ids)+')' if coast_ids else '')+')')
            args.extend([station,*alternatives,*coast_ids])
        if coast:
            coast_ids=list(dict.fromkeys([str(coast),*self.map_coast(coast)]))
            where.append('c.coast_id IN ('+','.join('?' for _ in coast_ids)+')');args.extend(coast_ids)
        if time_token:
            where.append('r.time_raw=?');args.append(time_token)
        if start or end:
            where.append("r.time_status IN ('explicit','assumed')" if allow_assumed else "r.time_status='explicit'")
            for column,op,value in [('valid_time','>=',start),('valid_time','<=',end)]:
                if value:
                    parsed=datetime.fromisoformat(str(value).replace('Z','+00:00'))
                    if parsed.tzinfo is None:raise ValueError('日期查询需要显式时区')
                    where.append('r.'+column+op+'?');args.append(parsed.astimezone(timezone.utc).isoformat())
        query=''' FROM numeric_cells v JOIN numeric_sheets s USING(sheet_id)
            JOIN numeric_files f USING(file_id)
            JOIN numeric_rows r ON r.sheet_id=v.sheet_id AND r.row_index=v.row_index
            JOIN numeric_columns c ON c.sheet_id=v.sheet_id AND c.col_index=v.col_index'''
        query+=(' WHERE '+' AND '.join(where)) if where else ''
        with self.connect() as db:
            total=db.execute('SELECT count(*)'+query,args).fetchone()[0]
            rows=db.execute('''SELECT f.file_id,f.source_id,f.filename,f.sha256,s.name AS sheet,s.unit,s.unit_status,
                s.vertical_datum,s.datum_status,s.ensemble_statistic,v.row_index,v.col_index,v.value,
                v.raw,v.quality,v.formula,c.label AS station,c.station_id,c.mapping_status,c.candidates,
                c.coast_id,c.coast_mapping_status,c.coast_candidates,
                r.time_raw,r.valid_time,r.time_status'''+query+' ORDER BY f.filename,s.sheet_id,v.row_index,v.col_index LIMIT ?',[*args,limit]).fetchall()
        items=[]
        for row in rows:
            item=dict(row);item['cell']=get_column_letter(item.pop('col_index'))+str(item.pop('row_index'))
            item['raw']=json.loads(item['raw']);item['candidates']=json.loads(item['candidates'])
            item['coast_candidates']=json.loads(item['coast_candidates'] or '[]')
            item['location_kind']='coast' if item['coast_id'] else 'station' if item['station_id'] else 'unresolved'
            items.append(item)
        return {'items':items,'total':total,'limit':limit,'scope':'原文件单元格数值，不代表单位、日期、基准和业务用途已核定',
                'time_query_note':'完整日期仅匹配有显式日期时区的记录；assumed须显式启用。缺月年的原时间可用time_token查询。'}

    def recover_coasts(self, geojson_path, coast_sql, *, source_id=None, expected_sha256=None, profile=None):
        self._writable()
        from forecast_runtime.storage import read_sql,coordinates
        geojson_path=Path(geojson_path);sha=digest(geojson_path)
        if expected_sha256 and sha!=expected_sha256:raise ValueError('岸段原件SHA256与来源登记不一致')
        data=json.loads(geojson_path.read_text('utf-8-sig'));features=data.get('features',[])
        sql=read_sql(coast_sql,'m_warn_tide_info_town');profile=profile or {}
        unit=str(profile.get('threshold_unit') or 'unknown')
        gathered={};issues=[];aliases=[]
        with self.connect() as db:
            db.execute('DELETE FROM numeric_coasts');db.execute('DELETE FROM numeric_thresholds');db.execute('DELETE FROM numeric_coast_conflicts')
            for version,records in [('geojson',features),('sql',sql)]:
                seen=set();gathered[version]={}
                for index,record in enumerate(records):
                    p=record.get('properties',{}) if version=='geojson' else record
                    coast_id=str(p.get('cid') if version=='geojson' else p.get('Id'))
                    if coast_id in seen:raise ValueError('岸段版本含重复ID，不能首条覆盖')
                    seen.add(coast_id)
                    name=str(p.get('ADMC') or p.get('TOWN') or '') if version=='geojson' else str(p.get('name') or '')
                    reference=p.get('CWZ') if version=='geojson' else p.get('StationName')
                    county=p.get('COUNTY') if version=='geojson' else p.get('County')
                    town=p.get('TOWN') if version=='geojson' else p.get('name')
                    for label in [name,str(county or '')+'-'+name]:
                        if label:aliases.append((coast_key(label),coast_id,'geojson_admc' if version=='geojson' else 'sql'))
                    for label in [town,str(county or '')+'-'+str(town or ''),str(county or '')+'-'+str(town or '')+'岸段']:
                        if label:aliases.append((coast_key(label),coast_id,'geojson_town' if version=='geojson' else 'sql'))
                    geometry=record.get('geometry') if version=='geojson' else None
                    sid,state,candidates=self.map_station(reference)
                    meta={'city':p.get('CITY') if version=='geojson' else p.get('Prefecture'),
                          'county':p.get('COUNTY') if version=='geojson' else p.get('County'),
                          'crs':data.get('crs') if version=='geojson' else None,'crs_status':'unknown',
                          'vertical_datum':'unknown','validity_period':'unknown',
                          'station_id':sid,'mapping_status':state,'station_candidates':candidates}
                    if version=='geojson':
                        meta.update(geometry_metadata(geometry))
                        if meta['coordinate_issues']:
                            issues.append({'version':version,'coast_id':coast_id,'name':name,'kind':'invalid_geometry_coordinates'})
                    if version=='sql':
                        lon,lat,coordinate_state=coordinates(p)
                        meta.update(lon=lon,lat=lat,coordinate_status=coordinate_state,
                                    formula=p.get('Formula'),base85=p.get('Base85'),is_verify=p.get('IsVerify'),
                                    observed_station_id=p.get('ObsOriStid'),forecast_station=p.get('FcstStation'),
                                    station_description_unit_hint='cm' if 'cm' in str(reference).lower() else None)
                    values={};known=[]
                    for colour in COLOURS:
                        raw=p.get(colour.upper()) if version=='geojson' else p.get(LEVELS[colour])
                        value,quality=number(raw);values[colour]={'raw':_json_value(raw),'value':value,'quality':quality}
                        if value is not None:known.append(value)
                        db.execute('INSERT INTO numeric_thresholds VALUES(?,?,?,?,?,?,?,?)',
                            (version,coast_id,colour,dump(_json_value(raw)),value,quality,unit,'assumed' if unit!='unknown' else 'unknown'))
                    if any(a>=b for a,b in zip(known,known[1:])):
                        issues.append({'version':version,'coast_id':coast_id,'name':name,'kind':'non_increasing_thresholds'})
                    meta['threshold_order']='strictly_increasing_known_levels' if len(known)>1 and all(a<b for a,b in zip(known,known[1:])) else 'insufficient_or_invalid'
                    db.execute('INSERT INTO numeric_coasts VALUES(?,?,?,?,?,?)',
                        (version,coast_id,name,reference,dump(geometry),dump(meta)))
                    gathered[version][coast_id]={'name':town,'county':county,'thresholds':values}
            for coast_id in sorted(set(gathered['geojson'])|set(gathered['sql'])):
                g=gathered['geojson'].get(coast_id);s=gathered['sql'].get(coast_id)
                if not g or not s:
                    db.execute('INSERT INTO numeric_coast_conflicts VALUES(?,?,?,?)',(coast_id,'presence',dump(bool(g)),dump(bool(s))))
                    continue
                for colour in COLOURS:
                    if (g['thresholds'][colour]['value'],g['thresholds'][colour]['quality'])!=(s['thresholds'][colour]['value'],s['thresholds'][colour]['quality']):
                        db.execute('INSERT INTO numeric_coast_conflicts VALUES(?,?,?,?)',
                            (coast_id,colour,dump(g['thresholds'][colour]),dump(s['thresholds'][colour])))
                for field in ('county','name'):
                    if name_key(g[field])!=name_key(s[field]):
                        db.execute('INSERT INTO numeric_coast_conflicts VALUES(?,?,?,?)',(coast_id,field,dump(g[field]),dump(s[field])))
            for role,path,metadata in [('coastal_geojson',geojson_path,{'feature_count':len(features),'source_id':source_id}),
                                       ('coastal_sql',Path(coast_sql),{'row_count':len(sql)})]:
                db.execute('INSERT OR REPLACE INTO numeric_reference_sources VALUES(?,?,?,?)',(role,path.name,digest(path),dump(metadata)))
            db.execute('DELETE FROM numeric_coast_aliases')
            db.executemany('INSERT OR IGNORE INTO numeric_coast_aliases VALUES(?,?,?)',aliases)
        self.coast_aliases=None
        mapped=self.refresh_coast_mapping()
        return {'geojson_features':len(features),'sql_rows':len(sql),'issues':issues,'conflicts':self.coast_conflicts(),'column_coast_mapping':mapped}

    def map_coast(self,label):
        if self.coast_aliases is None:
            self.coast_aliases={}
            with self.connect() as db:
                rows=db.execute('SELECT alias,coast_id,source_version FROM numeric_coast_aliases').fetchall()
            for row in rows:self.coast_aliases.setdefault(row['alias'],[]).append(row)
            self.coast_aliases={k:_coast_candidates(v) for k,v in self.coast_aliases.items()}
        return self.coast_aliases.get(coast_key(label),[])

    def refresh_coast_mapping(self):
        self._writable()
        counts={}
        with self.connect() as db:
            for row in db.execute('SELECT DISTINCT label FROM numeric_columns').fetchall():
                label=row[0]
                candidates=_coast_candidates(db.execute('SELECT coast_id,source_version FROM numeric_coast_aliases WHERE alias=?',(coast_key(label),)).fetchall())
                status='unique_source_coast_name' if len(candidates)==1 else 'ambiguous' if candidates else 'unmapped'
                n=db.execute('UPDATE numeric_columns SET coast_id=?,coast_mapping_status=?,coast_candidates=? WHERE label=?',
                    (candidates[0] if len(candidates)==1 else None,status,dump(candidates),label)).rowcount
                counts[status]=counts.get(status,0)+n
        return counts

    def coast_conflicts(self):
        with self.connect() as db:
            rows=db.execute('SELECT * FROM numeric_coast_conflicts ORDER BY coast_id,field').fetchall()
        result=[]
        for row in rows:
            g=json.loads(row['geojson_value']);s=json.loads(row['sql_value'])
            kind=('presence' if row['field']=='presence' else 'identity_conflict' if row['field'] not in COLOURS else
                  'missing_encoding_difference' if g.get('value') is None and s.get('value') is None else 'threshold_value_conflict')
            result.append({**dict(row),'geojson_value':g,'sql_value':s,'kind':kind})
        return result

    def coasts(self, query='', version=None):
        where=[];args=[]
        if query:
            ids=list(dict.fromkeys([str(query),*self.map_coast(query)]))
            where.append('coast_id IN ('+','.join('?' for _ in ids)+')');args.extend(ids)
        if version:
            if version not in {'geojson','sql'}:raise ValueError('岸段版本仅支持geojson或sql')
            where.append('version=?');args.append(version)
        with self.connect() as db:
            rows=db.execute('SELECT * FROM numeric_coasts'+(' WHERE '+' AND '.join(where) if where else '')+' ORDER BY coast_id,version',args).fetchall()
            items=[]
            for row in rows:
                item=dict(row);item['metadata']=json.loads(item['metadata']);item['geometry']=json.loads(item['geometry'])
                thresholds=db.execute('SELECT colour,raw,value,quality,unit,unit_status FROM numeric_thresholds WHERE version=? AND coast_id=? ORDER BY colour',(item['version'],item['coast_id'])).fetchall()
                item['thresholds']=[{**dict(t),'raw':json.loads(t['raw'])} for t in thresholds];items.append(item)
        return {'items':items,'conflicts':[r for r in self.coast_conflicts() if not query or r['coast_id'] in {item['coast_id'] for item in items}],
                'scope':'两个源版本并列存档；未自动选择阈值现行版本，invalid_sentinel不参与阈值计算。'}

    def summary(self):
        conflicts=self.coast_conflicts()
        with self.connect() as db:
            counts=lambda table:db.execute('SELECT count(*) FROM '+table).fetchone()[0]
            return {'parser_version':PARSER_VERSION,'files':counts('numeric_files'),'sheets':counts('numeric_sheets'),
                'rows':counts('numeric_rows'),'cells':counts('numeric_cells'),
                'cell_quality':dict(db.execute('SELECT quality,count(*) FROM numeric_cells GROUP BY quality')),
                'time_status':dict(db.execute('SELECT time_status,count(*) FROM numeric_rows GROUP BY time_status')),
                'station_column_mapping':dict(db.execute('SELECT mapping_status,count(*) FROM numeric_columns WHERE coast_id IS NULL GROUP BY mapping_status')),
                'coast_column_mapping':dict(db.execute('SELECT coast_mapping_status,count(*) FROM numeric_columns GROUP BY coast_mapping_status')),
                'station_registry_rows':counts('numeric_stations'),'coast_versions':dict(db.execute('SELECT version,count(*) FROM numeric_coasts GROUP BY version')),
                'threshold_quality':dict(db.execute('SELECT quality,count(*) FROM numeric_thresholds GROUP BY quality')),
                'conflict_fields':counts('numeric_coast_conflicts'),'conflict_coasts':db.execute('SELECT count(distinct coast_id) FROM numeric_coast_conflicts').fetchone()[0],
                'threshold_value_conflict_fields':sum(c['kind']=='threshold_value_conflict' for c in conflicts),
                'threshold_value_conflict_coasts':len({c['coast_id'] for c in conflicts if c['kind']=='threshold_value_conflict'}),
                'missing_encoding_differences':sum(c['kind']=='missing_encoding_difference' for c in conflicts),
                'coast_identity_conflicts':sum(c['kind']=='identity_conflict' for c in conflicts),
                'business_interpretation':'staging_only_unknown_or_explicit_assumptions_not_official_validation',
                'pending_questions':['工作簿每列单位与高程基准','缺月年时区的时间标签如何绑定真实有效时刻','集合表的成员/分位数或统计含义','同名不同站及未匹配站点对应关系','岸段阈值现行版本、单位、高程基准和有效期']}

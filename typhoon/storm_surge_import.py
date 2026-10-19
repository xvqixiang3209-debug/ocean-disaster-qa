"""Import the 2001-2016 Zhejiang storm-surge workbooks into the typhoon SQLite KB."""

from __future__ import annotations

import argparse
import csv
import io
import json
import math
import os
import re
import shutil
import sqlite3
import subprocess
import tempfile
import zipfile
from datetime import datetime, timedelta, timezone
from pathlib import Path
from statistics import quantiles

from .config import (
    DEFAULT_DATABASE,
    DEFAULT_IBTRACS_ZIP,
    DEFAULT_STORM_NUMBER_MAP,
    DEFAULT_STORM_SURGE_ARCHIVE,
)


CHINA_STANDARD_TIME = timezone(timedelta(hours=8))


STATION_NAMES = {
    "AOJIANG": "鳌江",
    "BEIHAI": "北海",
    "BEILUN": "北仑",
    "BOAO": "博鳌",
    "CHIWANH": "赤湾港",
    "RUIAN": "瑞安",
    "RUIANH": "瑞安港",
    "WENZHOU1": "温州一站",
    "WENZHOU2": "温州二站",
    "KANMEN": "坎门",
    "DACHEN": "大陈",
    "DAISHAN": "岱山",
    "HAIMENZ": "海门",
    "JIANTIAO": "健跳",
    "JIAOJIAN": "椒江",
    "JIAOJIANG": "椒江",
    "ZHENHAIS": "镇海站",
    "ZHENHAIH": "镇海港",
    "GANPU": "澉浦",
    "ZHAPU": "乍浦",
    "JINSHAN": "金山",
    "DINGHAI": "定海",
    "DAJIESAN": "大戢山",
    "LUCHAOG": "芦潮港",
    "DONGTOU": "洞头",
    "FANGCG": "防城港",
    "GAOQIAO": "高桥",
    "GUANGZHO": "广州",
    "LONGGANG": "龙港",
    "LONGWAN": "龙湾",
    "NANJI": "南麂",
    "NINGBO": "宁波",
    "OUJIANGK": "瓯江口",
    "SANMEN": "三门",
    "SHACHENG": "沙埕",
    "SHANTOU": "汕头",
    "SHANTOUH": "汕头港",
    "SHANWEI": "汕尾",
    "SHENGSHAN": "嵊山",
    "SHENJIAM": "沈家门",
    "SHENJIAMEN": "沈家门",
    "SHIPU": "石浦",
    "SONGLANSH": "松兰山",
    "WUSONG": "吴淞",
    "ZHAPO": "闸坡",
}


STORM_SURGE_SCHEMA = """
DROP TABLE IF EXISTS storm_surge_observations;
DROP TABLE IF EXISTS storm_surge_stations;
DROP TABLE IF EXISTS storm_surge_events;

CREATE TABLE storm_surge_events (
    event_id TEXT PRIMARY KEY,
    storm_code TEXT NOT NULL UNIQUE,
    season INTEGER NOT NULL,
    storm_number INTEGER NOT NULL,
    sid TEXT,
    storm_name TEXT,
    link_method TEXT NOT NULL,
    link_confidence TEXT NOT NULL,
    time_match_status TEXT NOT NULL,
    time_offset_hours REAL,
    event_start TEXT,
    event_end TEXT,
    max_surge_m REAL,
    min_surge_m REAL,
    max_water_level_m REAL,
    peak_station_code TEXT,
    peak_station_name TEXT,
    peak_surge_time TEXT,
    severity TEXT,
    severity_basis TEXT,
    severity_lower_m REAL,
    severity_upper_m REAL,
    station_count INTEGER NOT NULL,
    observation_count INTEGER NOT NULL,
    qc_mismatch_count INTEGER NOT NULL,
    source_archive TEXT NOT NULL,
    source_file TEXT NOT NULL,
    FOREIGN KEY (sid) REFERENCES storms(sid)
);

CREATE TABLE storm_surge_stations (
    event_id TEXT NOT NULL,
    station_code TEXT NOT NULL,
    station_name TEXT NOT NULL,
    latitude REAL,
    longitude REAL,
    max_surge_m REAL,
    min_surge_m REAL,
    max_water_level_m REAL,
    peak_surge_time TEXT,
    peak_water_time TEXT,
    observation_count INTEGER NOT NULL,
    qc_mismatch_count INTEGER NOT NULL,
    PRIMARY KEY (event_id, station_code),
    FOREIGN KEY (event_id) REFERENCES storm_surge_events(event_id)
);

CREATE TABLE storm_surge_observations (
    event_id TEXT NOT NULL,
    station_code TEXT NOT NULL,
    observation_time TEXT NOT NULL,
    observed_water_level_m REAL NOT NULL,
    predicted_tide_m REAL NOT NULL,
    surge_m REAL NOT NULL,
    source_residual_m REAL,
    PRIMARY KEY (event_id, station_code, observation_time),
    FOREIGN KEY (event_id, station_code)
        REFERENCES storm_surge_stations(event_id, station_code)
);

CREATE INDEX idx_surge_events_sid ON storm_surge_events(sid);
CREATE INDEX idx_surge_events_season_number ON storm_surge_events(season, storm_number);
CREATE INDEX idx_surge_stations_event ON storm_surge_stations(event_id);
CREATE INDEX idx_surge_observations_time ON storm_surge_observations(observation_time);
"""


def _number(value) -> float | None:
    if value in (None, ""):
        return None
    try:
        number = float(value)
    except (TypeError, ValueError):
        return None
    if not math.isfinite(number) or abs(number) >= 5000 or number == -999:
        return None
    return number


def _iso(value: datetime | None) -> str | None:
    """Serialize source workbook times as China Standard Time (UTC+8)."""
    if value is None:
        return None
    localized = (
        value.replace(tzinfo=CHINA_STANDARD_TIME)
        if value.tzinfo is None
        else value.astimezone(CHINA_STANDARD_TIME)
    )
    return localized.isoformat(timespec="seconds")


def _parse_iso(value: str | None) -> datetime | None:
    """Parse an IBTrACS ISO time and compare it in China Standard Time."""
    if not value:
        return None
    try:
        parsed = datetime.fromisoformat(str(value).replace("Z", "+00:00"))
        if parsed.tzinfo is None:
            parsed = parsed.replace(tzinfo=timezone.utc)
        return parsed.astimezone(CHINA_STANDARD_TIME).replace(tzinfo=None)
    except ValueError:
        return None


def _event_date(value, year: int) -> datetime | None:
    number = _number(value)
    if number is None:
        return None
    code = int(round(number))
    month, day = divmod(code, 100)
    try:
        return datetime(year, month, day)
    except ValueError:
        return None


def _coordinate(header: str, label: str) -> float | None:
    match = re.search(rf"{label}\.\s*(\d{{2,3}})(?:\.(\d{{1,2}}))?'?", header, re.IGNORECASE)
    if not match:
        return None
    degrees_text, minutes_text = match.groups()
    if minutes_text is None and len(degrees_text) >= 4:
        minutes_text = degrees_text[-2:]
        degrees_text = degrees_text[:-2]
    degrees = int(degrees_text)
    minutes = int(minutes_text or 0)
    if minutes >= 60:
        return None
    return round(degrees + minutes / 60.0, 5)


def _sheet_rows(path: Path):
    if path.suffix.lower() == ".xlsx":
        from openpyxl import load_workbook

        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            for sheet in workbook.worksheets:
                yield sheet.title, [list(row) for row in sheet.iter_rows(values_only=True)]
        finally:
            workbook.close()
    else:
        import xlrd

        workbook = xlrd.open_workbook(path, on_demand=True)
        try:
            for sheet in workbook.sheets():
                yield sheet.name, [sheet.row_values(index) for index in range(sheet.nrows)]
        finally:
            workbook.release_resources()


def _cell(row: list, index: int):
    return row[index] if index < len(row) else None


def _refresh_station(station: dict, allowed_dates: set[datetime] | None = None) -> dict | None:
    dates = station["event_dates"]
    observations = station["observations"]
    if allowed_dates is not None:
        allowed_days = {item.date() for item in allowed_dates}
        dates = [item for item in dates if item.date() in allowed_days]
        observations = [item for item in observations if item["time"].date() in allowed_days]
    observations_by_time = {}
    for item in observations:
        observations_by_time.setdefault(item["time"], item)
    observations = list(observations_by_time.values())
    if not dates:
        return None
    peak_surge = max(observations, key=lambda item: item["surge_m"]) if observations else None
    minimum_surge = min(observations, key=lambda item: item["surge_m"]) if observations else None
    peak_water = max(observations, key=lambda item: item["observed_water_level_m"]) if observations else None
    station.update(
        {
            "event_dates": dates,
            "date_start": min(dates),
            "date_end": max(dates) + timedelta(hours=23),
            "max_surge_m": peak_surge["surge_m"] if peak_surge else None,
            "min_surge_m": minimum_surge["surge_m"] if minimum_surge else None,
            "max_water_level_m": peak_water["observed_water_level_m"] if peak_water else None,
            "peak_surge_time": peak_surge["time"] if peak_surge else None,
            "peak_water_time": peak_water["time"] if peak_water else None,
            "observation_count": len(observations),
            "observations": observations,
        }
    )
    return station


def _parse_station(path: Path, year: int, sheet_name: str, rows: list[list]) -> dict | None:
    if len(rows) < 5:
        return None
    station_code = re.sub(r"\s+", "", str(sheet_name or "")).upper()
    if not station_code:
        return None
    header = str(_cell(rows[0], 0) or "")
    observations = []
    event_dates = []
    qc_mismatches = 0
    for row_index in range(1, len(rows) - 1):
        predicted_row = rows[row_index]
        day = _event_date(_cell(predicted_row, 0), year)
        if day is None:
            continue
        event_dates.append(day)
        observed_row = rows[row_index - 1]
        residual_row = rows[row_index + 1]
        for hour in range(24):
            column = hour + 1
            observed_cm = _number(_cell(observed_row, column))
            predicted_cm = _number(_cell(predicted_row, column))
            if observed_cm is None or predicted_cm is None:
                continue
            surge_cm = observed_cm - predicted_cm
            source_residual_cm = _number(_cell(residual_row, column))
            if source_residual_cm is not None and abs(source_residual_cm - surge_cm) > 1.1:
                qc_mismatches += 1
                continue
            if abs(surge_cm) > 500:
                qc_mismatches += 1
                continue
            observations.append(
                {
                    "time": day + timedelta(hours=hour),
                    "observed_water_level_m": round(observed_cm / 100.0, 4),
                    "predicted_tide_m": round(predicted_cm / 100.0, 4),
                    "surge_m": round(surge_cm / 100.0, 4),
                    "source_residual_m": (
                        round(source_residual_cm / 100.0, 4) if source_residual_cm is not None else None
                    ),
                }
            )
    if not event_dates or (not observations and station_code.startswith("SHEET")):
        return None
    station = {
        "station_code": station_code,
        "station_name": STATION_NAMES.get(station_code, station_code),
        "latitude": _coordinate(header, "LAT"),
        "longitude": _coordinate(header, "LON"),
        "event_dates": event_dates,
        "qc_mismatch_count": qc_mismatches,
        "observations": observations,
    }
    return _refresh_station(station)


def _dominant_date_cluster(stations: list[dict]) -> set[datetime]:
    date_counts: dict[datetime, int] = {}
    for station in stations:
        for day in station["event_dates"]:
            date_counts[day] = date_counts.get(day, 0) + 1
    ordered = sorted(date_counts)
    if not ordered:
        return set()
    clusters: list[list[datetime]] = [[ordered[0]]]
    for day in ordered[1:]:
        if day - clusters[-1][-1] <= timedelta(days=3):
            clusters[-1].append(day)
        else:
            clusters.append([day])
    selected = max(clusters, key=lambda cluster: (sum(date_counts[day] for day in cluster), len(cluster)))
    return set(selected)


def _parse_workbook(path: Path) -> dict | None:
    match = re.fullmatch(r"(\d{2})(\d{2})", path.stem)
    if not match:
        return None
    year = 2000 + int(match.group(1))
    storm_number = int(match.group(2))
    stations = []
    declared_storm_codes: set[str] = set()
    for sheet_name, rows in _sheet_rows(path):
        header = str(_cell(rows[0], 0) or "") if rows else ""
        declared_storm_codes.update(
            match.group(1)
            for match in re.finditer(r"\bTP\s*[:：]?\s*(\d{4})\b", header, re.IGNORECASE)
        )
        station = _parse_station(path, year, sheet_name, rows)
        if station:
            stations.append(station)
    if not stations:
        return None
    selected_dates = _dominant_date_cluster(stations)
    stations = [refreshed for station in stations if (refreshed := _refresh_station(station, selected_dates))]
    if not stations:
        return None
    observations = [item for station in stations for item in station["observations"]]
    measured_stations = [station for station in stations if station["max_surge_m"] is not None]
    peak = max(measured_stations, key=lambda item: item["max_surge_m"]) if measured_stations else None
    return {
        "event_id": path.stem,
        "storm_code": path.stem,
        "season": year,
        "storm_number": storm_number,
        "event_start": min(item["date_start"] for item in stations),
        "event_end": max(item["date_end"] for item in stations),
        "max_surge_m": peak["max_surge_m"] if peak else None,
        "min_surge_m": min(
            (item["min_surge_m"] for item in measured_stations), default=None
        ),
        "max_water_level_m": max(
            (item["max_water_level_m"] for item in measured_stations), default=None
        ),
        "peak_station_code": peak["station_code"] if peak else None,
        "peak_station_name": peak["station_name"] if peak else None,
        "peak_surge_time": peak["peak_surge_time"] if peak else None,
        "station_count": len(stations),
        "observation_count": len(observations),
        "qc_mismatch_count": sum(item["qc_mismatch_count"] for item in stations),
        "source_file": path.name,
        "declared_storm_codes": sorted(declared_storm_codes),
        "stations": stations,
    }


def _ibtracs_number_sequence(
    source_zip: Path,
) -> tuple[dict[tuple[int, int], str], dict[str, datetime]]:
    first_times: dict[str, tuple[int, datetime]] = {}
    with zipfile.ZipFile(source_zip) as archive:
        members = [name for name in archive.namelist() if name.endswith("ibtracs.ALL.list.v04r01.csv")]
        if not members:
            return {}
        with archive.open(members[0]) as raw:
            stream = io.TextIOWrapper(raw, encoding="utf-8-sig", newline="")
            for row in csv.DictReader(stream):
                if str(row.get("BASIN") or "").strip() != "WP":
                    continue
                if str(row.get("TRACK_TYPE") or "main").strip().lower() not in {"", "main"}:
                    continue
                try:
                    season = int(row.get("SEASON") or 0)
                    category = float(row.get("TOKYO_GRADE") or -99)
                except (TypeError, ValueError):
                    continue
                if category < 2:
                    continue
                sid = str(row.get("SID") or "").strip()
                timestamp = _parse_iso(str(row.get("ISO_TIME") or "").replace(" ", "T") + "Z")
                if not sid or timestamp is None:
                    continue
                previous = first_times.get(sid)
                if previous is None or timestamp < previous[1]:
                    first_times[sid] = (season, timestamp)
    by_season: dict[int, list[tuple[datetime, str]]] = {}
    for sid, (season, timestamp) in first_times.items():
        by_season.setdefault(season, []).append((timestamp, sid))
    result = {}
    for season, items in by_season.items():
        for number, (_, sid) in enumerate(sorted(items), start=1):
            result[(season, number)] = sid
    return result, {sid: timestamp for sid, (_, timestamp) in first_times.items()}


def _load_storms(database_path: Path) -> dict[str, dict]:
    with sqlite3.connect(database_path) as connection:
        connection.row_factory = sqlite3.Row
        rows = connection.execute(
            """
            SELECT sid, season, name, start_time, end_time, closest_time_zhejiang,
                   closest_distance_zhejiang, affects_zhejiang
            FROM storms
            """
        ).fetchall()
        chinese_names = {
            str(row["sid"]): str(row["storm_name_cn"] or "").strip()
            for row in connection.execute(
                """SELECT sid, MAX(storm_name_cn) AS storm_name_cn
                   FROM typhoon_report_cases
                   WHERE sid IS NOT NULL AND TRIM(COALESCE(storm_name_cn, '')) <> ''
                   GROUP BY sid"""
            )
        }
    storms = {row["sid"]: dict(row) for row in rows}
    for sid, storm in storms.items():
        storm["name_cn"] = chinese_names.get(sid, "")
    return storms


def _time_gap_hours(event_start: datetime, event_end: datetime, storm: dict) -> float | None:
    """Return zero for overlap, otherwise the gap between two time intervals."""
    storm_start = _parse_iso(storm.get("start_time"))
    storm_end = _parse_iso(storm.get("end_time"))
    if storm_start is None or storm_end is None:
        return None
    if storm_start <= event_end and storm_end >= event_start:
        return 0.0
    if storm_start > event_end:
        return (storm_start - event_end).total_seconds() / 3600.0
    return (event_start - storm_end).total_seconds() / 3600.0


def _time_status(gap_hours: float | None) -> str:
    if gap_hours is None:
        return "unavailable"
    if gap_hours == 0:
        return "overlap"
    if gap_hours <= 24:
        return "nearby"
    return "conflict"


def _linked_result(storm: dict, method: str, confidence: str, gap_hours: float) -> dict:
    return {
        "sid": storm["sid"],
        "storm_name": storm.get("name"),
        "storm_name_cn": storm.get("name_cn"),
        "link_method": method,
        "link_confidence": confidence,
        "time_match_status": _time_status(gap_hours),
        "time_offset_hours": round(gap_hours, 2),
    }


def _link_event(
    event: dict,
    storms: dict[str, dict],
    official_mapping: dict[str, dict],
    number_sequence: dict[tuple[int, int], str],
    designation_times: dict[str, datetime],
) -> dict:
    event_start = event["event_start"]
    event_end = event["event_end"]
    midpoint = event_start + (event_end - event_start) / 2

    candidates = []
    for storm in storms.values():
        if storm.get("season") != event["season"]:
            continue
        start = _parse_iso(storm.get("start_time"))
        end = _parse_iso(storm.get("end_time"))
        closest = _parse_iso(storm.get("closest_time_zhejiang"))
        if start is None or end is None or closest is None:
            continue
        gap_hours = _time_gap_hours(event_start, event_end, storm)
        if gap_hours is None or gap_hours > 72:
            continue
        time_difference = abs((closest - midpoint).total_seconds()) / 3600.0
        distance = float(storm.get("closest_distance_zhejiang") or 9999)
        score = gap_hours * 4.0 + time_difference + min(distance, 3000.0) / 30.0
        candidates.append((score, gap_hours, distance, storm))
    candidates.sort(key=lambda item: item[0])
    temporal = candidates[0] if candidates else None

    official = official_mapping.get(event["storm_code"])
    official_storm = storms.get(official.get("sid")) if official else None
    if official_storm:
        official_gap = _time_gap_hours(event_start, event_end, official_storm)
        expected_name = str(official.get("name") or "").strip().upper()
        actual_name = str(official_storm.get("name") or "").strip().upper()
        name_matches = not expected_name or expected_name == actual_name
        if name_matches and official_gap == 0:
            return _linked_result(
                official_storm, "JMA官方编号+名称+事件时间", "high", official_gap
            )
        if name_matches and official_gap is not None and official_gap <= 24:
            # A different storm that overlaps the gauge window and passes much closer
            # to Zhejiang is stronger evidence than a near-but-nonoverlapping code match.
            if temporal:
                _, temporal_gap, temporal_distance, temporal_storm = temporal
                official_distance = float(official_storm.get("closest_distance_zhejiang") or 9999)
                if (
                    temporal_storm["sid"] != official_storm["sid"]
                    and temporal_gap == 0
                    and temporal_distance + 300 < official_distance
                ):
                    return _linked_result(
                        temporal_storm,
                        "事件时间+浙江路径（官方编号时间冲突）",
                        "medium",
                        temporal_gap,
                    )
            return _linked_result(
                official_storm, "JMA官方编号+名称+相邻事件时间", "medium", official_gap
            )
        if temporal:
            _, temporal_gap, _, temporal_storm = temporal
            return _linked_result(
                temporal_storm,
                "事件时间+浙江路径（官方编号时间冲突）",
                "medium",
                temporal_gap,
            )

    sequence_sid = number_sequence.get((event["season"], event["storm_number"]))
    sequence_storm = storms.get(sequence_sid) if sequence_sid else None
    if sequence_storm:
        sequence_gap = _time_gap_hours(event_start, event_end, sequence_storm)
        if sequence_gap is not None and sequence_gap <= 24:
            confidence = "high" if sequence_gap == 0 else "medium"
            return _linked_result(
                sequence_storm, "台风年编号+事件时间", confidence, sequence_gap
            )
    if temporal:
        _, temporal_gap, _, temporal_storm = temporal
        return _linked_result(temporal_storm, "事件时间+浙江路径", "medium", temporal_gap)
    _ = designation_times  # Retained for compatibility with older callers.
    return {
        "sid": None,
        "storm_name": None,
        "link_method": "未匹配（无72小时内候选台风）",
        "link_confidence": "none",
        "time_match_status": "conflict",
        "time_offset_hours": None,
    }


def _assign_severity(events: list[dict]) -> tuple[float | None, float | None]:
    values = sorted(event["max_surge_m"] for event in events if event.get("max_surge_m") is not None)
    if len(values) < 3:
        lower, upper = 0.5, 1.5
    else:
        lower, upper = quantiles(values, n=3, method="inclusive")
    for event in events:
        value = event.get("max_surge_m")
        if value is None:
            event["severity"] = None
            event["severity_basis"] = "最大增水缺失"
        elif value <= lower:
            event["severity"] = "轻微"
            event["severity_basis"] = f"最大增水位于2001—2016样本低三分位（≤{lower:.2f} m）"
        elif value >= upper:
            event["severity"] = "严重"
            event["severity_basis"] = f"最大增水位于2001—2016样本高三分位（≥{upper:.2f} m）"
        else:
            event["severity"] = "一般"
            event["severity_basis"] = "最大增水位于2001—2016样本中间三分位"
        event["severity_lower_m"] = lower
        event["severity_upper_m"] = upper
    return lower, upper


def _load_official_mapping() -> dict[str, dict]:
    number_map_path = Path(DEFAULT_STORM_NUMBER_MAP)
    if not number_map_path.exists():
        return {}
    raw_mapping = json.loads(number_map_path.read_text(encoding="utf-8"))
    return raw_mapping.get("mapping", raw_mapping)


def prepare_storm_surge_workbook(
    filename: str,
    data: bytes,
    database_path: str | Path,
    *,
    relative_path: str | None = None,
) -> dict | None:
    """Recognize, parse and deterministically link one uploaded surge workbook.

    Recognition is intentionally strict: the filename must be a four-digit
    JMA annual number known by the packaged mapping, and at least one worksheet
    header must declare the same ``TP`` number.  Unrelated spreadsheets can
    therefore continue through the normal RAG importer without being mistaken
    for structured storm-surge observations.
    """

    safe_name = Path(filename).name
    extension = Path(safe_name).suffix.lower()
    code = Path(safe_name).stem
    if extension not in {".xls", ".xlsx"} or not re.fullmatch(r"\d{4}", code):
        return None
    official_mapping = _load_official_mapping()
    if code not in official_mapping:
        return None
    database_path = Path(database_path).resolve()
    if not database_path.exists():
        raise FileNotFoundError(f"台风知识库不存在: {database_path}")

    with tempfile.TemporaryDirectory(prefix=f"storm_surge_{code}_") as directory:
        temp_path = Path(directory) / safe_name
        temp_path.write_bytes(data)
        event = _parse_workbook(temp_path)
    if not event:
        return None
    declared_codes = set(event.get("declared_storm_codes") or [])
    if code not in declared_codes:
        return None

    storms = _load_storms(database_path)
    event.update(_link_event(event, storms, official_mapping, {}, {}))
    _assign_severity([event])
    event["source_file"] = relative_path or safe_name
    return event


def storm_surge_event_summary(event: dict) -> dict:
    """Return the JSON-safe association summary exposed by the upload API."""

    return {
        "recognized": True,
        "storm_code": event["storm_code"],
        "sid": event.get("sid"),
        "storm_name": event.get("storm_name"),
        "storm_name_cn": event.get("storm_name_cn"),
        "season": event["season"],
        "link_method": event["link_method"],
        "link_confidence": event["link_confidence"],
        "time_match_status": event["time_match_status"],
        "time_offset_hours": event.get("time_offset_hours"),
        "event_start": _iso(event.get("event_start")),
        "event_end": _iso(event.get("event_end")),
        "station_count": event["station_count"],
        "observation_count": event["observation_count"],
        "qc_mismatch_count": event["qc_mismatch_count"],
        "max_surge_m": event.get("max_surge_m"),
        "peak_station_code": event.get("peak_station_code"),
        "peak_station_name": event.get("peak_station_name"),
    }


def storm_surge_rag_summary(event: dict) -> str:
    """Build a compact semantic preface so derived maxima are retrievable by name."""

    def metric(value) -> str:
        return "未记录" if value is None else f"{value} m"

    names = [str(event.get("storm_name_cn") or "").strip(), str(event.get("storm_name") or "").strip()]
    name_text = " / ".join(name for name in names if name) or "名称未提供"
    lines = [
        "台风风暴潮结构化摘要（由上传工作簿逐站逐时数据自动计算）",
        f"台风：{event['season']}年 {name_text}；年编号：{event['storm_code']}；SID：{event.get('sid') or '未关联'}。",
        f"事件时间：{_iso(event.get('event_start'))} 至 {_iso(event.get('event_end'))}。",
        (
            f"全事件最大增水：{metric(event.get('max_surge_m'))}；"
            f"最大增水站点：{event.get('peak_station_name') or '未提供'}"
            f"（{event.get('peak_station_code') or '未提供'}）；"
            f"最大增水时间：{_iso(event.get('peak_surge_time')) or '未提供'}。"
        ),
        f"共{event['station_count']}个站点、{event['observation_count']}条有效逐时观测。",
        "各站统计：",
    ]
    for station in sorted(event["stations"], key=lambda item: item["station_code"]):
        lines.append(
            f"- {station['station_name']}（{station['station_code']}）："
            f"最大增水 {metric(station.get('max_surge_m'))}，"
            f"最高水位 {metric(station.get('max_water_level_m'))}，"
            f"最大增水时间 {_iso(station.get('peak_surge_time')) or '未提供'}，"
            f"有效观测 {station.get('observation_count', 0)} 条。"
        )
    lines.append(
        f"自动关联：{event.get('link_method')}；置信度 {event.get('link_confidence')}；"
        f"事件时间校验 {event.get('time_match_status')}。"
    )
    return "\n".join(lines)


def _extract_archive(archive_path: Path, output_dir: Path) -> list[Path]:
    seven_zip = shutil.which("7z") or shutil.which("7zz")
    if not seven_zip:
        raise RuntimeError("未找到7z，无法解压RAR风暴潮数据")
    subprocess.run(
        [seven_zip, "x", "-y", f"-o{output_dir}", str(archive_path)],
        check=True,
        capture_output=True,
        text=True,
    )
    return sorted(
        path for path in output_dir.rglob("*")
        if path.is_file() and path.suffix.lower() in {".xls", ".xlsx"}
    )


def _insert_event(connection: sqlite3.Connection, event: dict, source_label: str) -> None:
    connection.execute(
        """
        INSERT INTO storm_surge_events (
            event_id, storm_code, season, storm_number, sid, storm_name,
            link_method, link_confidence, time_match_status, time_offset_hours,
            event_start, event_end,
            max_surge_m, min_surge_m, max_water_level_m,
            peak_station_code, peak_station_name, peak_surge_time,
            severity, severity_basis, severity_lower_m, severity_upper_m,
            station_count, observation_count, qc_mismatch_count,
            source_archive, source_file
        ) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)
        """,
        (
            event["event_id"], event["storm_code"], event["season"], event["storm_number"],
            event.get("sid"), event.get("storm_name"), event["link_method"], event["link_confidence"],
            event["time_match_status"], event["time_offset_hours"],
            _iso(event["event_start"]), _iso(event["event_end"]), event["max_surge_m"],
            event["min_surge_m"], event["max_water_level_m"], event["peak_station_code"],
            event["peak_station_name"], _iso(event["peak_surge_time"]), event["severity"],
            event["severity_basis"], event["severity_lower_m"], event["severity_upper_m"],
            event["station_count"], event["observation_count"], event["qc_mismatch_count"],
            source_label, event["source_file"],
        ),
    )
    for station in event["stations"]:
        connection.execute(
            """
            INSERT INTO storm_surge_stations (
                event_id, station_code, station_name, latitude, longitude,
                max_surge_m, min_surge_m, max_water_level_m,
                peak_surge_time, peak_water_time, observation_count, qc_mismatch_count
            ) VALUES (?,?,?,?,?,?,?,?,?,?,?,?)
            """,
            (
                event["event_id"], station["station_code"], station["station_name"],
                station["latitude"], station["longitude"], station["max_surge_m"],
                station["min_surge_m"], station["max_water_level_m"],
                _iso(station["peak_surge_time"]), _iso(station["peak_water_time"]),
                station["observation_count"], station["qc_mismatch_count"],
            ),
        )
        connection.executemany(
            """
            INSERT INTO storm_surge_observations (
                event_id, station_code, observation_time,
                observed_water_level_m, predicted_tide_m, surge_m, source_residual_m
            ) VALUES (?,?,?,?,?,?,?)
            """,
            [
                (
                    event["event_id"], station["station_code"], _iso(item["time"]),
                    item["observed_water_level_m"], item["predicted_tide_m"],
                    item["surge_m"], item["source_residual_m"],
                )
                for item in station["observations"]
            ],
        )


def _write_database(database_path: Path, archive_path: Path, events: list[dict]) -> None:
    temp_handle = tempfile.NamedTemporaryFile(
        prefix="typhoon_with_surge_", suffix=".sqlite3", dir=database_path.parent, delete=False
    )
    temp_path = Path(temp_handle.name)
    temp_handle.close()
    shutil.copy2(database_path, temp_path)
    try:
        connection = sqlite3.connect(temp_path)
        connection.execute("PRAGMA foreign_keys=ON")
        connection.executescript(STORM_SURGE_SCHEMA)
        for event in events:
            _insert_event(connection, event, str(archive_path))
        metadata = {
            "storm_surge_source": str(archive_path),
            "storm_surge_imported_at": datetime.now(timezone.utc).isoformat(),
            "storm_surge_event_count": len(events),
            "storm_surge_station_count": sum(event["station_count"] for event in events),
            "storm_surge_observation_count": sum(event["observation_count"] for event in events),
            "storm_surge_linked_event_count": sum(bool(event.get("sid")) for event in events),
            "storm_surge_time_overlap_count": sum(
                event.get("time_match_status") == "overlap" for event in events
            ),
            "storm_surge_time_nearby_count": sum(
                event.get("time_match_status") == "nearby" for event in events
            ),
            "storm_surge_time_conflict_count": sum(
                event.get("time_match_status") == "conflict" for event in events
            ),
            "storm_surge_schema_version": 2,
        }
        connection.executemany(
            "INSERT OR REPLACE INTO metadata(key, value) VALUES (?, ?)",
            [(key, json.dumps(value, ensure_ascii=False)) for key, value in metadata.items()],
        )
        connection.commit()
        connection.execute("ANALYZE")
        connection.close()
        os.replace(temp_path, database_path)
    except Exception:
        try:
            connection.close()
        except Exception:
            pass
        if temp_path.exists():
            temp_path.unlink()
        raise


def merge_storm_surge_events(
    database_path: str | Path,
    events: list[dict],
    *,
    source_label: str = "visual_knowledge_upload",
) -> dict:
    """Transactionally upsert uploaded surge events without deleting other years."""

    database_path = Path(database_path).resolve()
    if not events:
        return {
            "event_count": 0,
            "linked_event_count": 0,
            "station_count": 0,
            "observation_count": 0,
            "high_confidence_links": 0,
        }
    event_by_code = {event["storm_code"]: event for event in events}
    events = list(event_by_code.values())
    connection = sqlite3.connect(database_path, timeout=30)
    try:
        connection.execute("PRAGMA foreign_keys=ON")
        table_exists = connection.execute(
            "SELECT 1 FROM sqlite_master WHERE type='table' AND name='storm_surge_events'"
        ).fetchone()
        if not table_exists:
            connection.executescript(STORM_SURGE_SCHEMA)

        placeholders = ",".join("?" for _ in events)
        replaced_ids = [event["event_id"] for event in events]
        existing_values = [
            float(row[0])
            for row in connection.execute(
                f"SELECT max_surge_m FROM storm_surge_events "
                f"WHERE event_id NOT IN ({placeholders}) AND max_surge_m IS NOT NULL",
                replaced_ids,
            )
        ]
        combined_values = existing_values + [
            float(event["max_surge_m"])
            for event in events
            if event.get("max_surge_m") is not None
        ]
        if len(combined_values) < 3:
            lower, upper = 0.5, 1.5
        else:
            lower, upper = quantiles(sorted(combined_values), n=3, method="inclusive")
        for event in events:
            value = event.get("max_surge_m")
            if value is None:
                event.update(severity=None, severity_basis="最大增水缺失")
            elif value <= lower:
                event.update(
                    severity="轻微",
                    severity_basis=f"最大增水位于2001—2016样本低三分位（≤{lower:.2f} m）",
                )
            elif value >= upper:
                event.update(
                    severity="严重",
                    severity_basis=f"最大增水位于2001—2016样本高三分位（≥{upper:.2f} m）",
                )
            else:
                event.update(
                    severity="一般",
                    severity_basis="最大增水位于2001—2016样本中间三分位",
                )
            event["severity_lower_m"] = lower
            event["severity_upper_m"] = upper

        connection.execute("BEGIN IMMEDIATE")
        for event in events:
            event_id = event["event_id"]
            connection.execute("DELETE FROM storm_surge_observations WHERE event_id = ?", (event_id,))
            connection.execute("DELETE FROM storm_surge_stations WHERE event_id = ?", (event_id,))
            connection.execute("DELETE FROM storm_surge_events WHERE event_id = ?", (event_id,))
            _insert_event(connection, event, source_label)

        all_rows = connection.execute(
            "SELECT event_id, max_surge_m FROM storm_surge_events"
        ).fetchall()
        for event_id, value in all_rows:
            if value is None:
                severity, basis = None, "最大增水缺失"
            elif value <= lower:
                severity, basis = "轻微", f"最大增水位于2001—2016样本低三分位（≤{lower:.2f} m）"
            elif value >= upper:
                severity, basis = "严重", f"最大增水位于2001—2016样本高三分位（≥{upper:.2f} m）"
            else:
                severity, basis = "一般", "最大增水位于2001—2016样本中间三分位"
            connection.execute(
                """UPDATE storm_surge_events
                   SET severity=?, severity_basis=?, severity_lower_m=?, severity_upper_m=?
                   WHERE event_id=?""",
                (severity, basis, lower, upper, event_id),
            )

        totals = connection.execute(
            """SELECT COUNT(*), COALESCE(SUM(station_count),0),
                      COALESCE(SUM(observation_count),0),
                      COALESCE(SUM(CASE WHEN sid IS NOT NULL THEN 1 ELSE 0 END),0)
               FROM storm_surge_events"""
        ).fetchone()
        metadata = {
            "storm_surge_source": source_label,
            "storm_surge_imported_at": datetime.now(timezone.utc).isoformat(),
            "storm_surge_event_count": int(totals[0]),
            "storm_surge_station_count": int(totals[1]),
            "storm_surge_observation_count": int(totals[2]),
            "storm_surge_linked_event_count": int(totals[3]),
            "storm_surge_last_batch_event_count": len(events),
            "storm_surge_schema_version": 2,
        }
        connection.executemany(
            "INSERT OR REPLACE INTO metadata(key, value) VALUES (?, ?)",
            [(key, json.dumps(value, ensure_ascii=False)) for key, value in metadata.items()],
        )
        connection.commit()
    except Exception:
        connection.rollback()
        raise
    finally:
        connection.close()
    return {
        "event_count": len(events),
        "linked_event_count": sum(bool(event.get("sid")) for event in events),
        "station_count": sum(event["station_count"] for event in events),
        "observation_count": sum(event["observation_count"] for event in events),
        "high_confidence_links": sum(event.get("link_confidence") == "high" for event in events),
        "time_overlap_links": sum(event.get("time_match_status") == "overlap" for event in events),
        "database_event_count": int(totals[0]),
        "severity_lower_m": lower,
        "severity_upper_m": upper,
    }


def import_storm_surge_archive(
    archive_path: str | Path = DEFAULT_STORM_SURGE_ARCHIVE,
    database_path: str | Path = DEFAULT_DATABASE,
    ibtracs_zip: str | Path = DEFAULT_IBTRACS_ZIP,
    *,
    dry_run: bool = False,
    progress=None,
) -> dict:
    archive_path = Path(archive_path).resolve()
    database_path = Path(database_path).resolve()
    ibtracs_zip = Path(ibtracs_zip).resolve()
    if not archive_path.exists():
        raise FileNotFoundError(f"风暴潮压缩包不存在: {archive_path}")
    if not database_path.exists():
        raise FileNotFoundError(f"台风知识库不存在: {database_path}")
    if not ibtracs_zip.exists():
        raise FileNotFoundError(f"IBTrACS压缩包不存在: {ibtracs_zip}")

    storms = _load_storms(database_path)
    official_mapping = _load_official_mapping()
    number_sequence, designation_times = _ibtracs_number_sequence(ibtracs_zip)
    with tempfile.TemporaryDirectory(prefix="storm_surge_import_") as directory:
        workbook_paths = _extract_archive(archive_path, Path(directory))
        events = []
        for index, path in enumerate(workbook_paths, start=1):
            event = _parse_workbook(path)
            if event:
                event.update(
                    _link_event(event, storms, official_mapping, number_sequence, designation_times)
                )
                events.append(event)
            if progress:
                progress(index, len(workbook_paths), path.name)

    lower, upper = _assign_severity(events)
    if not dry_run:
        _write_database(database_path, archive_path, events)
    return {
        "archive": str(archive_path),
        "database": str(database_path),
        "workbook_count": len(workbook_paths),
        "event_count": len(events),
        "station_count": sum(event["station_count"] for event in events),
        "observation_count": sum(event["observation_count"] for event in events),
        "linked_event_count": sum(bool(event.get("sid")) for event in events),
        "high_confidence_links": sum(event.get("link_confidence") == "high" for event in events),
        "medium_confidence_links": sum(event.get("link_confidence") == "medium" for event in events),
        "low_confidence_links": sum(event.get("link_confidence") == "low" for event in events),
        "time_overlap_links": sum(event.get("time_match_status") == "overlap" for event in events),
        "time_nearby_links": sum(event.get("time_match_status") == "nearby" for event in events),
        "time_conflicts": sum(event.get("time_match_status") == "conflict" for event in events),
        "unlinked_events": [event["storm_code"] for event in events if not event.get("sid")],
        "severity_lower_m": lower,
        "severity_upper_m": upper,
        "dry_run": dry_run,
        "links": [
            {
                "storm_code": event["storm_code"],
                "sid": event.get("sid"),
                "name": event.get("storm_name"),
                "method": event["link_method"],
                "confidence": event["link_confidence"],
                "time_match_status": event["time_match_status"],
                "time_offset_hours": event["time_offset_hours"],
                "max_surge_m": event["max_surge_m"],
                "severity": event["severity"],
            }
            for event in events
        ],
    }


def main() -> None:
    parser = argparse.ArgumentParser(description="导入浙江历史台风风暴潮Excel数据")
    parser.add_argument("--source", default=str(DEFAULT_STORM_SURGE_ARCHIVE), help="RAR压缩包路径")
    parser.add_argument("--database", default=str(DEFAULT_DATABASE), help="台风SQLite知识库路径")
    parser.add_argument("--ibtracs", default=str(DEFAULT_IBTRACS_ZIP), help="IBTrACS压缩包路径")
    parser.add_argument("--dry-run", action="store_true", help="只解析和匹配，不写入数据库")
    args = parser.parse_args()

    def progress(index, total, filename):
        if index == total or index % 10 == 0:
            print(f"解析风暴潮文件 {index}/{total}: {filename}", flush=True)

    result = import_storm_surge_archive(
        args.source,
        args.database,
        args.ibtracs,
        dry_run=args.dry_run,
        progress=progress,
    )
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
